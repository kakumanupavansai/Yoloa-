"""HTTP API for YOLOA.

Assessment decisions live here (and not in React): eligibility, timing, scoring,
answer ownership, and completion are all calculated against persisted records.
"""
import calendar as cal
import base64
import hashlib
import hmac
import io
import json
import logging
import mimetypes
import os
import re
import secrets
import string
import unicodedata
import uuid
from datetime import date, datetime, timedelta
from decimal import Decimal
from urllib.parse import urlencode

import requests
from django.conf import settings
from django.contrib.auth import authenticate, login as django_login, logout as django_logout, update_session_auth_hash
from django.contrib.auth.password_validation import validate_password
from django.contrib.auth.tokens import default_token_generator
from django.core.exceptions import ValidationError as DjangoValidationError
from django.core.mail import send_mail
from django.core.validators import URLValidator, validate_email
from django.http import FileResponse, HttpResponseRedirect
from django.middleware.csrf import get_token
from django.db import IntegrityError, transaction
from django.db.models import Avg, Count, F, IntegerField, Max, OuterRef, Q, Subquery, Value
from django.db.models.functions import Coalesce
from django.shortcuts import get_object_or_404
from django.utils.encoding import force_bytes, force_str
from django.utils.http import urlsafe_base64_encode, urlsafe_base64_decode
from django.utils.text import get_valid_filename
from django.utils import timezone
from django.views.decorators.clickjacking import xframe_options_sameorigin
from django.views.decorators.csrf import ensure_csrf_cookie
from django.views.decorators.http import require_GET
from rest_framework.decorators import api_view, parser_classes, permission_classes
from rest_framework.authentication import SessionAuthentication
from rest_framework.exceptions import APIException, AuthenticationFailed, ValidationError
from rest_framework.parsers import FormParser, JSONParser, MultiPartParser
from rest_framework.permissions import AllowAny
from rest_framework.response import Response
from rest_framework_simplejwt.tokens import RefreshToken
from rest_framework_simplejwt.authentication import JWTAuthentication
from rest_framework_simplejwt.views import TokenRefreshView

from . import models as m
from .rate_limits import (RateLimitUnavailable, check_login_attempt, clear as clear_rate_limit,
                          hit as hit_rate_limit, request_ip)
from .utils import IsAdminRole, IsStudentRole, abs_url, to_int, user_json

ANY = [JSONParser, FormParser, MultiPartParser]
logger = logging.getLogger(__name__)
PASS_PERCENT = Decimal('60')
FINAL_DURATION_SECONDS = 600
MAX_IMPORT_BYTES = 5 * 1024 * 1024


def fail(message, code=400):
    return Response({'message': message}, status=code)


def value_number(value):
    """JSON-friendly number which does not expose Decimal implementation details."""
    return float(value) if isinstance(value, Decimal) else value


def date_filter(request, queryset, field='created_at'):
    """Apply optional ISO date range query parameters to an analytics queryset."""
    for name, lookup in (('start', 'gte'), ('end', 'lte')):
        raw = request.query_params.get(name)
        if not raw:
            continue
        try:
            parsed = date.fromisoformat(raw)
        except ValueError:
            raise ValidationError({name: 'Use YYYY-MM-DD.'})
        if lookup == 'gte':
            queryset = queryset.filter(**{f'{field}__date__gte': parsed})
        else:
            queryset = queryset.filter(**{f'{field}__date__lte': parsed})
    return queryset


# ───────────────────────── Authentication and dashboards ─────────────────────────
@ensure_csrf_cookie
@api_view(['GET'])
@permission_classes([AllowAny])
def csrf_token(request):
    return Response({'csrfToken': get_token(request)})


def user_for_email_login(identifier):
    """Accept email addresses while retaining username login for existing API clients."""
    if '@' not in identifier:
        return identifier
    users = m.User.objects.filter(email__iexact=identifier)
    return users.first().username if users.count() == 1 else ''


def enforce_csrf(request):
    """Apply Django's CSRF check to public DRF actions such as login and reset."""
    if isinstance(getattr(request, 'successful_authenticator', None), JWTAuthentication):
        return
    SessionAuthentication().enforce_csrf(request)


def login_rate_response(request, identifier):
    """Return a generic denial when either shared login threshold has been reached."""
    try:
        retry_after = check_login_attempt(request, identifier)
    except RateLimitUnavailable:
        return Response({'message': 'Sign-in is temporarily unavailable. Please try again shortly.'}, status=503)
    if retry_after:
        response = Response({'message': 'Too many sign-in attempts. Please wait before trying again.'}, status=429)
        response['Retry-After'] = str(retry_after)
        return response
    return None


@api_view(['POST'])
@permission_classes([AllowAny])
def login(request):
    enforce_csrf(request)
    data = request.data
    identifier = str(data.get('email') or data.get('username') or '').strip()
    password = data.get('password', '')
    limited = login_rate_response(request, identifier)
    if limited:
        return limited
    if not identifier or not password:
        return fail('Enter your email address and password.', 400)
    username = user_for_email_login(identifier)
    user = authenticate(request=request, username=username, password=password) if username else None
    if not user:
        return fail('Invalid email or password.', 401)
    # The submitted role only selects which portal is in use; privileges always come from the stored user.
    portal = data.get('role')
    if portal and portal != user.role:
        if portal == 'admin':
            return fail('This account does not have administrator access.', 403)
        if user.role == 'admin':
            return fail('This is an administrator account. Use administrator sign-in instead.', 403)
        return fail('This account cannot sign in here.', 403)
    django_login(request, user)
    m.LoginEvent.objects.create(user=user, event_type='login', auth_method='email_password')
    return Response({'user': user_json(user, request)})


class TrackedTokenRefreshView(TokenRefreshView):
    """Compatibility for the documented Postman integration; browser auth does not use refresh tokens."""
    def post(self, request, *args, **kwargs):
        try:
            refresh_token = RefreshToken(request.data.get('refresh', ''))
            user_id = refresh_token.get('user_id')
            user = m.User.objects.only('auth_token_version').filter(pk=user_id).first()
            if not user or int(refresh_token.get('av', 0)) != user.auth_token_version:
                raise AuthenticationFailed('Token has been revoked.')
        except Exception:
            return fail('Token is invalid or has been revoked.', 401)
        response = super().post(request, *args, **kwargs)
        if response.status_code < 400:
            try:
                token = RefreshToken(request.data.get('refresh', ''))
                user_id = token.get('user_id')
                if user_id and m.User.objects.filter(pk=user_id).exists():
                    m.LoginEvent.objects.create(user_id=user_id, event_type='refresh')
            except Exception:
                pass
        return response


@api_view(['POST'])
@permission_classes([AllowAny])
def legacy_jwt_login(request):
    """Temporary non-browser JWT login for the existing Postman collection."""
    identifier = str(request.data.get('email') or request.data.get('username') or '').strip()
    limited = login_rate_response(request, identifier)
    if limited:
        return limited
    username = user_for_email_login(identifier)
    user = authenticate(request=request, username=username, password=request.data.get('password', '')) if username else None
    if not user:
        return fail('Invalid username or password', 401)
    refresh = RefreshToken.for_user(user)
    refresh['av'] = user.auth_token_version
    m.LoginEvent.objects.create(user=user, event_type='login', auth_method='legacy_jwt')
    return Response({'token': str(refresh.access_token), 'refresh': str(refresh), 'user': user_json(user, request)})


@api_view(['POST'])
@permission_classes([AllowAny])
def register(request):
    enforce_csrf(request)
    data = request.data
    name = ' '.join(str(data.get('fullName', '')).split())
    email = str(data.get('email', '')).strip().lower()
    password = data.get('password', '')
    if not name or not email or '@' not in email:
        return fail('Enter your full name and a valid email address.', 400)
    if password != data.get('passwordConfirmation'):
        return fail('Passwords do not match.', 400)
    try:
        validate_password(password)
    except DjangoValidationError as error:
        return Response({'message': 'Please choose a stronger password.', 'errors': {'password': error.messages}}, status=400)
    with transaction.atomic():
        if m.User.objects.filter(email__iexact=email).exists():
            return fail('An account with this email already exists.', 409)
        stem = re.sub(r'[^a-z0-9._-]+', '', email.split('@', 1)[0].lower()) or 'student'
        username, suffix = stem, 1
        while m.User.objects.filter(username__iexact=username).exists():
            suffix += 1
            username = f'{stem}{suffix}'
        parts = name.split(' ', 1)
        user = m.User.objects.create_user(username=username, email=email, password=password,
                                          first_name=parts[0], last_name=parts[1] if len(parts) > 1 else '', role='student')
    django_login(request, user)
    m.LoginEvent.objects.create(user=user, event_type='login', auth_method='email_password')
    return Response({'user': user_json(user, request)}, status=201)


DEV_EMAIL_BACKENDS = {
    'django.core.mail.backends.console.EmailBackend',
    'django.core.mail.backends.filebased.EmailBackend',
    'django.core.mail.backends.locmem.EmailBackend',
}


def email_delivery_mode():
    """'smtp' when real delivery is configured, 'development' for console-style backends, else 'unavailable'."""
    backend = settings.EMAIL_BACKEND
    if backend in DEV_EMAIL_BACKENDS:
        return 'development'
    if backend == 'django.core.mail.backends.smtp.EmailBackend':
        return 'smtp' if settings.EMAIL_HOST else 'unavailable'
    if backend == 'django.core.mail.backends.dummy.EmailBackend':
        return 'unavailable'
    return 'smtp'


RESET_MESSAGES = {
    'smtp': 'If an account exists for that email, a password reset link has been sent. Check your inbox.',
    'development': ('If an account exists for that email, a password reset link has been generated. '
                    'Email delivery is in development mode, so the link is written to the server console instead of being emailed.'),
}


def reset_user(uid, token):
    """Return the active user for a reset link, or None when it is malformed, expired, or already used."""
    try:
        user = m.User.objects.get(pk=force_str(urlsafe_base64_decode(str(uid or ''))), is_active=True)
    except (ValueError, TypeError, OverflowError, m.User.DoesNotExist):
        return None
    # Django's token embeds the password hash and last login, so it stops working once used.
    return user if default_token_generator.check_token(user, str(token or '')) else None


@api_view(['POST'])
@permission_classes([AllowAny])
def password_reset_request(request):
    enforce_csrf(request)
    mode = email_delivery_mode()
    if mode == 'unavailable':
        # Checked before any account lookup, so this response never reveals whether an email is registered.
        return fail('Password reset email is not available right now. Please contact your administrator.', 503)
    email = str(request.data.get('email', '')).strip().lower()
    if not email or '@' not in email:
        return fail('Enter a valid email address.', 400)
    try:
        account_retry = hit_rate_limit('password-reset-account', email,
                                       settings.PASSWORD_RESET_ACCOUNT_LIMIT,
                                       settings.PASSWORD_RESET_RATE_WINDOW_SECONDS)
        source_retry = hit_rate_limit('password-reset-source', request_ip(request),
                                      settings.PASSWORD_RESET_SOURCE_LIMIT,
                                      settings.PASSWORD_RESET_RATE_WINDOW_SECONDS)
    except RateLimitUnavailable:
        return fail('Password reset is temporarily unavailable. Please try again shortly.', 503)
    if account_retry or source_retry:
        return fail('Please wait before requesting another reset email.', 429)
    user = m.User.objects.filter(email__iexact=email, is_active=True).first() if email else None
    if user:
        uid = urlsafe_base64_encode(force_bytes(user.pk))
        token = default_token_generator.make_token(user)
        reset_url = f'{settings.FRONTEND_URL}/reset-password?{urlencode({"uid": uid, "token": token})}'
        try:
            send_mail('Reset your YOLOA password',
                      f'Use this one-time link to choose a new password. It expires according to the account security policy:\n\n{reset_url}\n\nIf you did not request this, ignore this message.',
                      settings.DEFAULT_FROM_EMAIL, [user.email], fail_silently=False)
        except Exception as error:
            # Keep the public response identical; never log the link or token.
            logger.error('Password reset email delivery failed: %s', type(error).__name__)
    return Response({'message': RESET_MESSAGES[mode], 'delivery': mode})


@api_view(['POST'])
@permission_classes([AllowAny])
def password_reset_validate(request):
    enforce_csrf(request)
    if not reset_user(request.data.get('uid'), request.data.get('token')):
        return fail('This password reset link is invalid, has expired, or has already been used.', 400)
    return Response({'valid': True})


@api_view(['POST'])
@permission_classes([AllowAny])
def password_reset_confirm(request):
    enforce_csrf(request)
    user = reset_user(request.data.get('uid'), request.data.get('token'))
    password = request.data.get('password', '')
    if not user:
        return fail('This password reset link is invalid, has expired, or has already been used.', 400)
    confirmation = request.data.get('passwordConfirmation')
    if confirmation is not None and confirmation != password:
        return fail('Passwords do not match.', 400)
    try:
        validate_password(password, user)
    except DjangoValidationError as error:
        return Response({'message': 'Please choose a stronger password.', 'errors': {'password': error.messages}}, status=400)
    user.set_password(password)
    user.save(update_fields=['password'])
    m.User.objects.filter(pk=user.pk).update(auth_token_version=F('auth_token_version') + 1)
    return Response({'message': 'Password updated. Please sign in with your new password.'})


@api_view(['POST'])
def logout(request):
    enforce_csrf(request)
    if request.user.is_authenticated:
        m.User.objects.filter(pk=request.user.pk).update(auth_token_version=F('auth_token_version') + 1)
    django_logout(request)
    return Response({'message': 'Signed out.'})


def google_configured():
    return bool(settings.GOOGLE_OAUTH_CLIENT_ID and settings.GOOGLE_OAUTH_CLIENT_SECRET and settings.GOOGLE_OAUTH_REDIRECT_URI)


@api_view(['GET'])
@permission_classes([AllowAny])
def google_status(request):
    """Public, secret-free flag so the student sign-in page can explain a missing configuration up front."""
    return Response({'configured': google_configured()})


def frontend_auth_redirect(error=None, role=None):
    target = f'{settings.FRONTEND_URL}/login'
    if error:
        return HttpResponseRedirect(f'{target}?auth_error={error}')
    return HttpResponseRedirect(f'{settings.FRONTEND_URL}/{"admin" if role == "admin" else "student"}/dashboard')


@require_GET
def google_login_start(request):
    if not google_configured():
        return frontend_auth_redirect(error='google_not_configured')
    link_user_id = ''
    if request.GET.get('mode') == 'link':
        if not request.user.is_authenticated:
            return frontend_auth_redirect(error='sign_in_before_linking')
        if request.user.role != 'student':
            return frontend_auth_redirect(error='google_student_only')
        link_user_id = str(request.user.pk)
    state = secrets.token_urlsafe(32)
    nonce = secrets.token_urlsafe(32)
    verifier = secrets.token_urlsafe(64)
    request.session['google_oauth'] = {'state': state, 'nonce': nonce, 'verifier': verifier, 'link_user_id': link_user_id}
    challenge = base64.urlsafe_b64encode(hashlib.sha256(verifier.encode()).digest()).rstrip(b'=').decode()
    query = urlencode({'client_id': settings.GOOGLE_OAUTH_CLIENT_ID,
                       'redirect_uri': settings.GOOGLE_OAUTH_REDIRECT_URI,
                       'response_type': 'code', 'scope': 'openid email profile', 'state': state,
                       'nonce': nonce, 'code_challenge': challenge, 'code_challenge_method': 'S256',
                       'prompt': 'select_account'})
    return HttpResponseRedirect(f'https://accounts.google.com/o/oauth2/v2/auth?{query}')


def verify_google_id_token(encoded_token, client_id):
    from google.auth.transport.requests import Request as GoogleRequest
    from google.oauth2 import id_token
    return id_token.verify_oauth2_token(encoded_token, GoogleRequest(), client_id)


@require_GET
def google_login_callback(request):
    oauth = request.session.pop('google_oauth', None)
    state = request.GET.get('state', '')
    code = request.GET.get('code', '')
    if request.GET.get('error') or not oauth or not hmac.compare_digest(state, oauth.get('state', '')) or not code:
        return frontend_auth_redirect(error='google_cancelled_or_invalid')
    try:
        response = requests.post('https://oauth2.googleapis.com/token', data={
            'code': code, 'client_id': settings.GOOGLE_OAUTH_CLIENT_ID,
            'client_secret': settings.GOOGLE_OAUTH_CLIENT_SECRET,
            'redirect_uri': settings.GOOGLE_OAUTH_REDIRECT_URI,
            'grant_type': 'authorization_code', 'code_verifier': oauth['verifier'],
        }, timeout=10)
        response.raise_for_status()
        identity = verify_google_id_token(response.json()['id_token'], settings.GOOGLE_OAUTH_CLIENT_ID)
    except Exception:
        return frontend_auth_redirect(error='google_verification_failed')
    if not hmac.compare_digest(str(identity.get('nonce', '')), oauth['nonce']) or identity.get('email_verified') is not True:
        return frontend_auth_redirect(error='google_verification_failed')
    subject, email = str(identity.get('sub', '')), str(identity.get('email', '')).strip().lower()
    if not subject or not email:
        return frontend_auth_redirect(error='google_verification_failed')
    if oauth.get('link_user_id'):
        user = request.user
        if (not user.is_authenticated or user.role != 'student' or
                str(user.pk) != oauth['link_user_id'] or user.email.lower() != email):
            return frontend_auth_redirect(error='google_link_mismatch')
        if m.User.objects.exclude(pk=user.pk).filter(google_sub=subject).exists():
            return frontend_auth_redirect(error='google_account_already_linked')
        user.google_sub = subject
        user.save(update_fields=['google_sub'])
        return frontend_auth_redirect(role=user.role)
    user = m.User.objects.filter(google_sub=subject).first()
    if not user:
        if m.User.objects.filter(email__iexact=email).exists():
            return frontend_auth_redirect(error='google_account_link_required')
        parts = ' '.join(str(identity.get('name', '')).split()).split(' ', 1)
        stem = re.sub(r'[^a-z0-9._-]+', '', email.split('@', 1)[0]) or 'student'
        username, suffix = stem, 1
        while m.User.objects.filter(username__iexact=username).exists():
            suffix += 1
            username = f'{stem}{suffix}'
        user = m.User.objects.create_user(username=username, email=email, first_name=parts[0],
                                          last_name=parts[1] if len(parts) > 1 else '', role='student', google_sub=subject)
    # A previously linked administrator account must never become a Google-authenticated session.
    if user.role != 'student':
        return frontend_auth_redirect(error='google_student_only')
    django_login(request, user)
    m.LoginEvent.objects.create(user=user, event_type='login', auth_method='google')
    return frontend_auth_redirect(role=user.role)


@api_view(['GET'])
def me(request):
    return Response(user_json(request.user, request))


@api_view(['POST'])
def password_change(request):
    """Change the signed-in user's own password. The account is always request.user, never a submitted id."""
    enforce_csrf(request)
    user = request.user
    data = request.data
    current, new = data.get('currentPassword', ''), data.get('newPassword', '')
    if not user.has_usable_password():
        return fail('This account signs in with Google and has no password yet. Use “Forgot password” to set one.', 400)
    if not current or not new:
        return fail('Enter your current password and a new password.', 400)
    try:
        retry_after = hit_rate_limit('password-change', user.pk, settings.PASSWORD_CHANGE_RATE_LIMIT,
                                     settings.PASSWORD_CHANGE_RATE_WINDOW_SECONDS)
    except RateLimitUnavailable:
        return fail('Password change is temporarily unavailable. Please try again shortly.', 503)
    if retry_after:
        return fail('Too many incorrect attempts. Please wait before trying again.', 429)
    if not user.check_password(current):
        return Response({'message': 'Your current password is incorrect.',
                         'errors': {'currentPassword': ['Your current password is incorrect.']}}, status=400)
    if new != data.get('newPasswordConfirmation'):
        return Response({'message': 'The new passwords do not match.',
                         'errors': {'newPasswordConfirmation': ['The new passwords do not match.']}}, status=400)
    if new == current:
        return Response({'message': 'Choose a password different from your current one.',
                         'errors': {'newPassword': ['Choose a password different from your current one.']}}, status=400)
    try:
        validate_password(new, user)
    except DjangoValidationError as error:
        return Response({'message': 'Please choose a stronger password.', 'errors': {'newPassword': error.messages}}, status=400)
    try:
        clear_rate_limit('password-change', user.pk)
    except RateLimitUnavailable:
        return fail('Password change is temporarily unavailable. Please try again shortly.', 503)
    user.set_password(new)
    user.save(update_fields=['password'])
    m.User.objects.filter(pk=user.pk).update(auth_token_version=F('auth_token_version') + 1)
    # Keeps this session signed in; every other session of this account stops validating (its stored hash no longer matches).
    update_session_auth_hash(request, user)
    return Response({'message': 'Your password has been changed.'})


def iso(value):
    return value.isoformat() if value else None


# ───────────────────────── Module progression (single source of truth) ─────────────────────────
FINALIZED = ('submitted', 'expired')


def pass_mark(exam):
    """Finals always use 60%; module and legacy quizzes use their configured mark (default 50%)."""
    if exam.assessment_type == 'final_exam':
        return PASS_PERCENT
    return Decimal(exam.pass_percent if exam.pass_percent else 50)


def attempt_stats(student, exams):
    """Attempts, latest and best results per assessment id in two queries (no N+1)."""
    ids = [exam.id for exam in exams]
    stats = {exam_id: {'attempts': 0, 'latestScore': None, 'latestPercent': None, 'latestPassed': None, 'bestScore': None,
                       'bestPercent': None, 'passed': False, 'lastAttemptAt': None} for exam_id in ids}
    if not ids:
        return stats
    rows = m.Attempt.objects.filter(student=student, exam_id__in=ids, status__in=FINALIZED).order_by('exam_id', '-submitted_at', '-id')
    for attempt in rows:
        item = stats[attempt.exam_id]
        if item['attempts'] == 0:
            item.update(latestScore=value_number(attempt.score), latestPercent=value_number(attempt.percentage),
                        latestPassed=attempt.passed, lastAttemptAt=iso(attempt.submitted_at))
        item['attempts'] += 1
        if item['bestPercent'] is None or attempt.percentage > Decimal(str(item['bestPercent'])):
            item['bestPercent'] = value_number(attempt.percentage)
            item['bestScore'] = value_number(attempt.score)
        item['passed'] = item['passed'] or attempt.passed
    return stats


def course_progress(student, course, chapters=None):
    """Return ordered module states for one student.

    A module with a published module quiz is complete only once that quiz has a passing finalized attempt.
    A module without a quiz is complete when the student explicitly marks it complete.
    Module N+1 unlocks only when every earlier *required* module is complete. The same function feeds
    the API responses and the server-side guards, so the browser can never unlock a module on its own.
    """
    chapters = list(course.chapters.order_by('order', 'id')) if chapters is None else sorted(
        chapters, key=lambda chapter: (chapter.order, chapter.id))
    quizzes = {exam.module_id: exam for exam in m.Exam.objects.filter(
        course=course, assessment_type='module_quiz', status='published', module__isnull=False).prefetch_related('questions')}
    stats = attempt_stats(student, list(quizzes.values()))
    marked = {row.chapter_id: row for row in m.ChapterProgress.objects.filter(student=student, chapter__in=chapters, completed=True)}
    modules, blocked = [], False
    for index, chapter in enumerate(chapters):
        quiz = quizzes.get(chapter.id)
        quiz_stats = stats.get(quiz.id) if quiz else None
        completed = bool(quiz_stats and quiz_stats['passed']) if quiz else chapter.id in marked
        progress_row = marked.get(chapter.id)
        modules.append({
            'id': chapter.id, 'title': chapter.title, 'description': chapter.description, 'order': index + 1,
            'duration': chapter.duration, 'required': chapter.required, 'completed': completed,
            'completedAt': iso(progress_row.completed_at) if progress_row else None,
            'unlocked': not blocked, 'locked': blocked,
            'quiz': ({'id': quiz.id, 'title': quiz.title, 'passPercent': quiz.pass_percent,
                      'questionCount': len(quiz.questions.all()), **quiz_stats} if quiz else None),
        })
        if chapter.required and not completed:
            blocked = True
    done = sum(1 for module in modules if module['completed'])
    next_module = next((module for module in modules if module['unlocked'] and not module['completed']), None)
    return {'modules': modules, 'completedModules': done, 'totalModules': len(modules),
            'percent': round(done * 100 / len(modules)) if modules else 0,
            'nextModule': {'id': next_module['id'], 'title': next_module['title']} if next_module else None}


def module_state(student, chapter):
    progress = course_progress(student, chapter.course)
    return next(module for module in progress['modules'] if module['id'] == chapter.id), progress


def recalc_enrollment(enrollment, progress=None):
    """Persist the derived percent and next module. A passed final (completion) is never undone here."""
    progress = progress or course_progress(enrollment.student, enrollment.course)
    enrollment.percent = progress['percent']
    modules = progress['modules']
    enrollment.current_chapter = (progress['nextModule']['title'] if progress['nextModule']
                                  else (modules[-1]['title'] if modules else ''))
    if enrollment.status != 'completed':
        enrollment.status = 'ongoing'
    enrollment.save(update_fields=['percent', 'current_chapter', 'status'])
    return enrollment


def get_enrollment(user, course):
    return m.Enrollment.objects.filter(student=user, course=course).first() if user.role == 'student' else None


def module_assessments(course):
    return m.Exam.objects.filter(course=course, assessment_type='module_quiz', status='published').select_related('module')


def latest_and_best(student, exam):
    return attempt_stats(student, [exam])[exam.id]


def final_eligibility(student, exam):
    if exam.assessment_type != 'final_exam':
        return {'eligible': True, 'reasons': []}
    reasons = []
    if exam.status != 'published':
        reasons.append('The final exam is not published.')
    if not exam.course_id:
        reasons.append('The final exam is not associated with a course.')
        return {'eligible': False, 'reasons': reasons, 'requiredQuizzes': []}
    enrollment = get_enrollment(student, exam.course)
    if not enrollment:
        reasons.append('Enroll in this course before starting its final exam.')
    required_modules = list(exam.course.chapters.filter(required=True))
    quizzes = {quiz.module_id: quiz for quiz in module_assessments(exam.course) if quiz.module_id}
    passed = set(m.Attempt.objects.filter(student=student, exam__in=quizzes.values(), passed=True, status__in=FINALIZED)
                 .values_list('exam_id', flat=True))
    report = []
    for module in required_modules:
        quiz = quizzes.get(module.id)
        if not quiz:
            report.append({'moduleId': module.id, 'title': module.title, 'status': 'missing'})
            reasons.append(f'A published required module quiz is missing for {module.title}.')
        elif quiz.id in passed:
            report.append({'moduleId': module.id, 'quizId': quiz.id, 'title': module.title, 'status': 'passed'})
        else:
            report.append({'moduleId': module.id, 'quizId': quiz.id, 'title': module.title, 'status': 'not_passed'})
            reasons.append(f'Pass the {module.title} module quiz first.')
    return {'eligible': not reasons, 'reasons': reasons, 'requiredQuizzes': report,
            'enrollmentStatus': enrollment.status if enrollment else None}


def final_summary(student, course):
    final = m.Exam.objects.filter(course=course, assessment_type='final_exam', status='published').prefetch_related('questions').first()
    if not final:
        return None
    eligibility = final_eligibility(student, final)
    active = m.Attempt.objects.filter(student=student, exam=final, status='in_progress').first()
    return {'id': final.id, 'title': final.title, 'durationSeconds': FINAL_DURATION_SECONDS, 'passPercent': int(PASS_PERCENT),
            'questionCount': len(final.questions.all()), 'inProgress': bool(active),
            'activeDeadline': iso(active.deadline_at) if active else None, **eligibility, **latest_and_best(student, final)}


def course_assessment_summary(student, course):
    progress = course_progress(student, course)
    quizzes = [{'id': module['quiz']['id'], 'moduleId': module['id'], 'moduleTitle': module['title'], 'required': module['required'],
                'unlocked': module['unlocked'], **{k: v for k, v in module['quiz'].items() if k != 'id'}}
               for module in progress['modules'] if module['quiz']]
    return {'moduleQuizzes': quizzes, 'finalExam': final_summary(student, course)}


def student_course_report(student, enrollment):
    """Full per-course progress used by the student progress page and the admin student detail."""
    progress = course_progress(student, enrollment.course)
    return {'courseId': enrollment.course_id, 'course': enrollment.course.title, 'category': enrollment.course.category,
            'status': enrollment.status, 'percent': progress['percent'], 'completionSource': enrollment.completion_source,
            'completedAt': iso(enrollment.completed_at), 'completedModules': progress['completedModules'],
            'totalModules': progress['totalModules'], 'nextModule': progress['nextModule'], 'modules': progress['modules'],
            'finalExam': final_summary(student, enrollment.course)}


# ───────────────────────── Dashboards ─────────────────────────
def recent_activity(student, limit=8):
    attempts = m.Attempt.objects.filter(student=student, status__in=FINALIZED).select_related('exam', 'exam__course', 'exam__module').order_by('-submitted_at')[:limit]
    items = [{'type': 'final' if attempt.exam.assessment_type == 'final_exam' else 'quiz', 'id': f'attempt-{attempt.id}',
              'title': attempt.exam.title, 'courseTitle': attempt.course_snapshot.get('title') or (attempt.exam.course.title if attempt.exam.course_id else ''),
              'percent': value_number(attempt.percentage), 'passed': attempt.passed, 'status': attempt.status, 'at': iso(attempt.submitted_at)}
             for attempt in attempts]
    manual = m.ChapterProgress.objects.filter(student=student, completed=True, completed_at__isnull=False,
                                              chapter__assessments__isnull=True).select_related('chapter__course').order_by('-completed_at')[:limit]
    items += [{'type': 'module', 'id': f'module-{row.id}', 'title': row.chapter.title, 'courseTitle': row.chapter.course.title,
               'percent': None, 'passed': True, 'status': 'completed', 'at': iso(row.completed_at)} for row in manual]
    return sorted(items, key=lambda item: item['at'] or '', reverse=True)[:limit]


@api_view(['GET'])
@permission_classes([IsStudentRole])
def student_dashboard(request):
    user = request.user
    finalize_expired_attempts()
    latest = m.Announcement.objects.first()
    courses, best_percents, quiz_attempts, quizzes_passed = [], [], 0, 0
    for enrollment in user.enrollments.select_related('course').order_by('-id'):
        report = student_course_report(user, enrollment)
        for module in report['modules']:
            quiz = module['quiz']
            if quiz and quiz['attempts']:
                best_percents.append(quiz['bestPercent'])
                quiz_attempts += quiz['attempts']
                quizzes_passed += 1 if quiz['passed'] else 0
        courses.append({key: report[key] for key in ('courseId', 'course', 'category', 'status', 'percent', 'completedModules',
                                                    'totalModules', 'nextModule', 'finalExam', 'completionSource')})
    in_progress = [course for course in courses if course['status'] != 'completed' and course['nextModule']]
    in_progress.sort(key=lambda course: -course['percent'])
    nxt = in_progress[0] if in_progress else None
    return Response({
        'announcement': latest.text if latest else '',
        'stats': {'enrolled': len(courses), 'completed': sum(1 for course in courses if course['status'] == 'completed'),
                  'inProgress': sum(1 for course in courses if course['status'] != 'completed'),
                  'averageQuizPercent': round(sum(best_percents) / len(best_percents)) if best_percents else None,
                  'quizAttempts': quiz_attempts, 'quizzesPassed': quizzes_passed},
        'courses': courses,
        'next': ({'courseId': nxt['courseId'], 'courseTitle': nxt['course'], 'moduleId': nxt['nextModule']['id'],
                  'moduleTitle': nxt['nextModule']['title'], 'percent': nxt['percent']} if nxt else None),
        'recentActivity': recent_activity(user),
        # Retained so the existing "completion" consumers keep working.
        'completion': [{'id': course['courseId'], 'title': course['course'], 'chapter': (course['nextModule'] or {}).get('title', ''),
                        'percent': course['percent'], 'status': course['status']} for course in courses],
    })


@api_view(['GET'])
@permission_classes([IsStudentRole])
def student_progress(request):
    finalize_expired_attempts()
    reports = [student_course_report(request.user, enrollment)
               for enrollment in request.user.enrollments.select_related('course').order_by('-id')]
    return Response({'courses': reports})


@api_view(['GET'])
@permission_classes([IsAdminRole])
def admin_dashboard(request):
    finalize_expired_attempts()
    students = m.User.objects.filter(role='student')
    period = date_filter(request, students, 'date_joined')
    now = timezone.now()
    week_ago = now - timedelta(days=7)
    enrollments = m.Enrollment.objects.filter(student__role='student')
    finalized = m.Attempt.objects.filter(status__in=FINALIZED)
    finals = finalized.filter(exam__assessment_type='final_exam')
    logins = m.LoginEvent.objects.filter(event_type='login', user__role='student')
    announcements = [{'id': announcement.id, 'author': announcement.author.display_name, 'text': announcement.text,
                      'createdAt': iso(announcement.created_at), 'avatar': abs_url(request, announcement.author.avatar)}
                     for announcement in m.Announcement.objects.select_related('author')[:5]]
    course_rows = []
    for course in m.Course.objects.annotate(
            modules=Count('chapters', distinct=True), students=Count('enrollments', distinct=True),
            completions=Count('enrollments', filter=Q(enrollments__status='completed'), distinct=True)).order_by('-id')[:8]:
        average = course.enrollments.aggregate(value=Avg('percent'))['value']
        course_rows.append({'id': course.id, 'title': course.title, 'category': course.category, 'isPublished': course.is_published,
                            'modules': course.modules, 'students': course.students, 'completions': course.completions,
                            'averageProgress': round(average) if average is not None else None})
    recent_logins = [{'id': event.id, 'userId': event.user_id, 'name': event.user.display_name, 'authMethod': event.auth_method,
                      'createdAt': iso(event.created_at)} for event in logins.select_related('user')[:6]]
    recent_attempts = [{'id': attempt.id, 'studentId': attempt.student_id, 'student': attempt.student.display_name,
                        'title': attempt.exam.title, 'type': attempt.exam.assessment_type, 'percent': value_number(attempt.percentage),
                        'passed': attempt.passed, 'submittedAt': iso(attempt.submitted_at)}
                       for attempt in finalized.select_related('student', 'exam').order_by('-submitted_at')[:6]]
    average_progress = enrollments.aggregate(value=Avg('percent'))['value']
    return Response({'announcements': announcements, 'courses': course_rows, 'recentLogins': recent_logins,
                     'recentAttempts': recent_attempts, 'stats': {
        'total': students.count(),
        'newRegistrations': period.count(),
        'newThisWeek': students.filter(date_joined__gte=week_ago).count(),
        'enrolled': students.filter(enrollments__isnull=False).distinct().count(),
        'courses': m.Course.objects.count(),
        'publishedCourses': m.Course.objects.filter(is_published=True).count(),
        'modules': m.Chapter.objects.count(),
        'enrollments': enrollments.count(),
        'completedEnrollments': enrollments.filter(status='completed').count(),
        'averageProgress': round(average_progress) if average_progress is not None else None,
        'quizAttempts': finalized.exclude(exam__assessment_type='final_exam').count(),
        'finalAttempts': finals.count(),
        'finalPassed': finals.filter(passed=True).count(),
        'activeLast7Days': logins.filter(created_at__gte=week_ago).values('user_id').distinct().count(),
    }})


@api_view(['POST'])
@permission_classes([IsAdminRole])
def announcements(request):
    text = (request.data.get('text') or '').strip()
    if not text:
        raise ValidationError({'text': 'This field is required.'})
    if len(text) > 2000:
        raise ValidationError({'text': 'Keep announcements under 2000 characters.'})
    announcement = m.Announcement.objects.create(author=request.user, text=text)
    return Response({'id': announcement.id, 'text': announcement.text, 'author': request.user.display_name}, status=201)


# ───────────────────────── Courses, modules, resources and enrollment ─────────────────────────
def parse_bool(raw, default=False):
    if raw is None or raw == '':
        return default
    if isinstance(raw, bool):
        return raw
    return str(raw).strip().lower() in ('1', 'true', 'yes', 'on')


def visible_courses_for(user):
    queryset = m.Course.objects.all()
    if user.role == 'student':
        queryset = queryset.filter(Q(is_published=True) | Q(enrollments__student=user)).distinct()
    return queryset


def course_row(course, enrollment, is_student, request):
    return {'id': course.id, 'title': course.title, 'description': course.description, 'category': course.category,
            'level': course.level, 'isPublished': course.is_published, 'cover': abs_url(request, course.cover),
            'enrollment': course.enrollment_label,
            'status': enrollment.status if enrollment else ('available' if is_student else ('published' if course.is_published else 'draft')),
            'enrolled': bool(enrollment), 'percent': enrollment.percent if enrollment else 0, 'chapters': course.n, 'modules': course.n,
            'students': course.s, 'completionSource': enrollment.completion_source if enrollment else '',
            'currentModule': enrollment.current_chapter if enrollment else '',
            'professor': course.professor.display_name if course.professor_id else ''}


def validate_course_payload(data, partial=False):
    values = {}
    if 'title' in data or not partial:
        title = str(data.get('title') or '').strip()
        if len(title) < 3:
            raise ValidationError({'title': 'Use a course title of at least 3 characters.'})
        if len(title) > 200:
            raise ValidationError({'title': 'Keep the title under 200 characters.'})
        values['title'] = title
    for js, field, limit in (('description', 'description', 5000), ('category', 'category', 80)):
        if js in data:
            values[field] = str(data.get(js) or '').strip()[:limit]
    if 'category' in values and not values['category']:
        values['category'] = 'General'
    if 'level' in data:
        level = str(data.get('level') or 'Beginner')
        if level not in dict(m.Course.LEVELS):
            raise ValidationError({'level': 'Choose Beginner, Intermediate or Advanced.'})
        values['level'] = level
    if 'isPublished' in data:
        values['is_published'] = parse_bool(data.get('isPublished'), True)
    for js, field in (('startDate', 'start_date'), ('endDate', 'end_date')):
        if js in data:
            raw = data.get(js)
            try:
                values[field] = date.fromisoformat(raw) if raw else None
            except ValueError:
                raise ValidationError({js: 'Use YYYY-MM-DD.'})
    if values.get('start_date') and values.get('end_date') and values['end_date'] < values['start_date']:
        raise ValidationError({'endDate': 'The end date must be after the start date.'})
    return values


def validate_cover(upload):
    if upload and (upload.size > 5 * 1024 * 1024 or not str(upload.content_type or '').startswith('image/')):
        raise ValidationError({'cover': 'Use an image smaller than 5 MB.'})
    return upload


def validate_module_payload(data, partial=False):
    values = {}
    if 'title' in data or not partial:
        title = str(data.get('title') or '').strip()
        if not title:
            raise ValidationError({'title': 'Give the module a title.'})
        values['title'] = title[:200]
    if 'description' in data:
        values['description'] = str(data.get('description') or '').strip()[:5000]
    if 'duration' in data:
        values['duration'] = str(data.get('duration') or '').strip()[:50]
    if 'required' in data:
        values['required'] = parse_bool(data.get('required'), True)
    if 'videoUrl' in data:
        url = str(data.get('videoUrl') or '').strip()
        if url:
            validate_http_url(url, 'videoUrl')
        values['video_url'] = url
    return values


@api_view(['GET', 'POST'])
@parser_classes(ANY)
def courses(request):
    if request.method == 'GET':
        is_student = request.user.role == 'student'
        enrollments = {item.course_id: item for item in request.user.enrollments.all()} if is_student else {}
        queryset = visible_courses_for(request.user).select_related('professor').annotate(
            n=Count('chapters', distinct=True), s=Count('enrollments', distinct=True)).order_by('id')
        return Response([course_row(course, enrollments.get(course.id), is_student, request) for course in queryset])

    if request.user.role != 'admin':
        return fail('Admin access required.', 403)
    data = request.data
    values = validate_course_payload(data)
    try:
        chapters = json.loads(data.get('chapters') or '[]') if isinstance(data.get('chapters', '[]'), str) else data.get('chapters') or []
    except json.JSONDecodeError:
        raise ValidationError({'chapters': 'Must be a JSON array.'})
    if not isinstance(chapters, list):
        raise ValidationError({'chapters': 'Must be a JSON array.'})
    module_values = [validate_module_payload({'title': chapter.get('title') or f'Module {index + 1}',
                                              'videoUrl': chapter.get('videoUrl', ''), 'description': chapter.get('description', ''),
                                              'required': chapter.get('required', True), 'duration': chapter.get('duration', '')})
                     for index, chapter in enumerate(chapters) if isinstance(chapter, dict)]
    with transaction.atomic():
        course = m.Course.objects.create(cover=validate_cover(request.FILES.get('cover')) or None,
                                         created_by=request.user, professor=request.user, **values)
        for index, module in enumerate(module_values):
            m.Chapter.objects.create(course=course, order=index, **module)
    return Response({'id': course.id, 'title': course.title}, status=201)


@api_view(['POST', 'DELETE'])
@permission_classes([IsStudentRole])
def course_enroll(request, pk):
    course = get_object_or_404(m.Course, pk=pk)
    if request.method == 'DELETE':
        m.Enrollment.objects.filter(student=request.user, course=course).delete()
        return Response(status=204)
    existing = m.Enrollment.objects.filter(student=request.user, course=course).first()
    if not existing and not course.is_published:
        return fail('This course is not open for enrollment.', 404)
    enrollment, created = m.Enrollment.objects.get_or_create(student=request.user, course=course)
    recalc_enrollment(enrollment)
    return Response({'enrolled': True, 'percent': enrollment.percent, 'status': enrollment.status,
                     'currentChapter': enrollment.current_chapter}, status=201 if created else 200)


def resource_json(item, include_admin=False):
    data = {'id': item.id, 'moduleId': item.chapter_id, 'type': item.resource_type, 'title': item.title, 'text': item.text,
            'url': item.url, 'hasFile': bool(item.file), 'fileName': item.original_name or (os.path.basename(item.file.name) if item.file else ''),
            'filePath': f'/resources/{item.id}/file' if item.file else '', 'order': item.order}
    if include_admin:
        data['createdAt'] = iso(item.created_at)
    return data


def module_admin_json(chapter, quiz_map):
    quiz = quiz_map.get(chapter.id)
    return {'id': chapter.id, 'title': chapter.title, 'description': chapter.description, 'duration': chapter.duration,
            'required': chapter.required, 'order': chapter.order, 'videoUrl': chapter.video_url,
            'resources': [resource_json(item, True) for item in chapter.resources.all()],
            'quiz': ({'id': quiz.id, 'title': quiz.title, 'status': quiz.status, 'passPercent': quiz.pass_percent,
                      'questionCount': len(quiz.questions.all()), 'attempts': quiz.attempt_count} if quiz else None)}


@api_view(['GET', 'PUT', 'PATCH', 'DELETE'])
@parser_classes(ANY)
def course_detail(request, pk):
    course = get_object_or_404(m.Course, pk=pk)
    if request.method in ('PUT', 'PATCH'):
        if request.user.role != 'admin':
            return fail('Admin access required.', 403)
        values = validate_course_payload(request.data, partial=True)
        cover = validate_cover(request.FILES.get('cover'))
        with transaction.atomic():
            for field, value in values.items():
                setattr(course, field, value)
            if cover:
                if course.cover:
                    course.cover.delete(save=False)
                course.cover = cover
            elif parse_bool(request.data.get('removeCover')) and course.cover:
                course.cover.delete(save=False)
                course.cover = None
            course.save()
        return Response({'id': course.id, 'title': course.title, 'isPublished': course.is_published})
    if request.method == 'DELETE':
        if request.user.role != 'admin':
            return fail('Admin access required.', 403)
        if course.exams.filter(attempts__isnull=False).exists():
            return fail('This course has assessment history and cannot be deleted. Unpublish it instead.', 409)
        with transaction.atomic():
            if course.cover:
                course.cover.delete(save=False)
            for item in course.contents.all():
                if item.file:
                    item.file.delete(save=False)
            # Import batches are staging data and protect their assessment; they go with it.
            m.QuestionImportBatch.objects.filter(exam__course=course).delete()
            course.exams.all().delete()
            course.delete()
        return Response(status=204)
    enrollment = get_enrollment(request.user, course)
    if request.user.role == 'student' and not course.is_published and not enrollment:
        return fail('Course not found.', 404)
    professor = course.professor
    chapters = list(course.chapters.annotate(resource_count=Count('resources')).order_by('order', 'id'))
    progress = course_progress(request.user, course, chapters) if enrollment else None
    states = {module['id']: module for module in (progress['modules'] if progress else [])}
    published = course.exams.filter(status='published').select_related('module').prefetch_related('questions')
    course_resources = [resource_json(item) for item in course.contents.filter(chapter__isnull=True)] if (enrollment or request.user.role == 'admin') else []
    return Response({'id': course.id, 'title': course.title, 'adminTitle': course.title, 'description': course.description,
        'category': course.category, 'level': course.level, 'isPublished': course.is_published,
        'cover': abs_url(request, course.cover), 'enrolled': bool(enrollment),
        'percent': progress['percent'] if progress else 0, 'completedModules': progress['completedModules'] if progress else 0,
        'status': enrollment.status if enrollment else 'available', 'completionSource': enrollment.completion_source if enrollment else '',
        'professor': {'name': professor.display_name, 'title': 'Instructor', 'avatar': abs_url(request, professor.avatar)} if professor else None,
        'content': course_resources,
        'nextModule': progress['nextModule'] if progress else None,
        # Before enrolling, students see the outline only (titles); resources require enrollment.
        'chapters': [{'id': chapter.id, 'title': chapter.title, 'description': chapter.description, 'duration': chapter.duration,
                      'required': chapter.required, 'completed': states.get(chapter.id, {}).get('completed', False),
                      'unlocked': states.get(chapter.id, {}).get('unlocked', False) if enrollment else False,
                      'quiz': states.get(chapter.id, {}).get('quiz'),
                      'resourceCount': chapter.resource_count} for chapter in chapters],
        'publishedExams': [exam_row(exam) for exam in published],
        'assessment': course_assessment_summary(request.user, course) if enrollment else None})


@api_view(['GET'])
@permission_classes([IsAdminRole])
def admin_course_detail(request, pk):
    """Everything an administrator needs to manage one course, including answer-free quiz summaries."""
    course = get_object_or_404(m.Course, pk=pk)
    chapters = list(course.chapters.prefetch_related('resources').order_by('order', 'id'))
    quiz_map = {exam.module_id: exam for exam in course.exams.filter(assessment_type='module_quiz', module__isnull=False)
                .annotate(attempt_count=Count('attempts')).prefetch_related('questions')}
    final = course.exams.filter(assessment_type='final_exam').annotate(attempt_count=Count('attempts')).prefetch_related('questions').first()
    enrollments = []
    for enrollment in course.enrollments.select_related('student').order_by('-id')[:200]:
        enrollments.append({'studentId': enrollment.student_id, 'name': enrollment.student.display_name, 'email': enrollment.student.email,
                            'status': enrollment.status, 'percent': enrollment.percent, 'currentModule': enrollment.current_chapter,
                            'completedAt': iso(enrollment.completed_at)})
    return Response({'id': course.id, 'title': course.title, 'description': course.description, 'category': course.category,
                     'level': course.level, 'isPublished': course.is_published, 'cover': abs_url(request, course.cover),
                     'startDate': iso(course.start_date), 'endDate': iso(course.end_date),
                     'modules': [module_admin_json(chapter, quiz_map) for chapter in chapters],
                     'resources': [resource_json(item, True) for item in course.contents.filter(chapter__isnull=True)],
                     'finalExam': ({'id': final.id, 'title': final.title, 'status': final.status, 'questionCount': len(final.questions.all()),
                                    'attempts': final.attempt_count, 'passPercent': int(PASS_PERCENT),
                                    'durationSeconds': FINAL_DURATION_SECONDS} if final else None),
                     'enrollments': enrollments,
                     'stats': {'students': course.enrollments.count(), 'completed': course.enrollments.filter(status='completed').count()}})


@api_view(['POST'])
@permission_classes([IsAdminRole])
def course_modules(request, pk):
    course = get_object_or_404(m.Course, pk=pk)
    values = validate_module_payload(request.data)
    order = (course.chapters.aggregate(value=Max('order'))['value'] or 0) + 1 if course.chapters.exists() else 0
    chapter = m.Chapter.objects.create(course=course, order=order, **values)
    refresh_course_enrollments(course)
    return Response(module_admin_json(chapter, {}), status=201)


@api_view(['POST'])
@permission_classes([IsAdminRole])
def course_module_order(request, pk):
    course = get_object_or_404(m.Course, pk=pk)
    ids = request.data.get('moduleIds')
    current = list(course.chapters.values_list('id', flat=True))
    if not isinstance(ids, list) or sorted(to_int(value, -1) for value in ids) != sorted(current):
        raise ValidationError({'moduleIds': 'Send every module id of this course exactly once.'})
    with transaction.atomic():
        for index, module_id in enumerate(ids):
            m.Chapter.objects.filter(pk=to_int(module_id), course=course).update(order=index)
    refresh_course_enrollments(course)
    return Response({'moduleIds': [to_int(value) for value in ids]})


def refresh_course_enrollments(course):
    """Module structure changes alter percentages; keep stored enrollment summaries consistent."""
    for enrollment in course.enrollments.select_related('student', 'course'):
        recalc_enrollment(enrollment)


@api_view(['GET', 'PUT', 'PATCH', 'DELETE'])
def module_detail(request, pk):
    chapter = get_object_or_404(m.Chapter.objects.select_related('course'), pk=pk)
    if request.method in ('PUT', 'PATCH', 'DELETE'):
        if request.user.role != 'admin':
            return fail('Admin access required.', 403)
        if request.method == 'DELETE':
            quizzes = m.Exam.objects.filter(module=chapter)
            if m.Attempt.objects.filter(exam__in=quizzes).exists():
                return fail('This module has quiz attempts and cannot be deleted. Mark it optional instead.', 409)
            with transaction.atomic():
                for item in chapter.resources.all():
                    if item.file:
                        item.file.delete(save=False)
                m.QuestionImportBatch.objects.filter(exam__in=quizzes).delete()
                quizzes.delete()
                course = chapter.course
                chapter.delete()
            refresh_course_enrollments(course)
            return Response(status=204)
        values = validate_module_payload(request.data, partial=True)
        for field, value in values.items():
            setattr(chapter, field, value)
        chapter.save()
        refresh_course_enrollments(chapter.course)
        return Response(module_admin_json(chapter, {}))
    # Student / admin learning view of one module.
    course = chapter.course
    if request.user.role == 'student':
        if not get_enrollment(request.user, course):
            return fail('Enroll in this course to start learning.', 403)
        state, progress = module_state(request.user, chapter)
        if state['locked']:
            return fail('Complete the previous required module to unlock this one.', 403)
    else:
        progress = {'modules': [{'id': item.id, 'title': item.title}
                                for item in course.chapters.order_by('order', 'id')]}
        state = {'completed': False, 'unlocked': True, 'locked': False, 'quiz': None}
    ordered = progress['modules']
    position = next(index for index, module in enumerate(ordered) if module['id'] == chapter.id)
    neighbour = lambda index: ({'id': ordered[index]['id'], 'title': ordered[index]['title'],
                                'locked': ordered[index].get('locked', False)} if 0 <= index < len(ordered) else None)
    return Response({'id': chapter.id, 'title': chapter.title, 'description': chapter.description, 'duration': chapter.duration,
                     'required': chapter.required, 'videoUrl': chapter.video_url, 'order': position + 1, 'totalModules': len(ordered),
                     'courseId': course.id, 'courseTitle': course.title, 'completed': state['completed'], 'quiz': state['quiz'],
                     'resources': [resource_json(item) for item in chapter.resources.all()],
                     'previous': neighbour(position - 1), 'next': neighbour(position + 1),
                     'coursePercent': progress.get('percent', 0)})


@api_view(['GET'])
def course_player(request, pk):
    """Compatibility endpoint: the course outline with lock states for an enrolled student."""
    course = get_object_or_404(m.Course, pk=pk)
    enrollment = get_enrollment(request.user, course)
    if request.user.role == 'student' and not enrollment:
        return fail('Enroll in this course to start learning.', 403)
    progress = course_progress(request.user, course)
    current = next((module for module in progress['modules'] if module['unlocked'] and not module['completed']), None)
    return Response({'courseId': course.id, 'courseTitle': course.title, 'description': course.description,
                     'percent': progress['percent'], 'status': enrollment.status if enrollment else 'ongoing',
                     'currentModuleId': current['id'] if current else None, 'lessonTitle': current['title'] if current else '',
                     'modules': progress['modules']})


@api_view(['PUT'])
@permission_classes([IsStudentRole])
def chapter_progress(request, pk):
    """Explicit completion for modules that have no quiz. Quiz modules complete only by passing the quiz."""
    chapter = get_object_or_404(m.Chapter.objects.select_related('course'), pk=pk)
    enrollment = get_enrollment(request.user, chapter.course)
    if not enrollment:
        return fail('Enroll in this course first.', 403)
    state, _ = module_state(request.user, chapter)
    if state['locked']:
        return fail('Complete the previous required module to unlock this one.', 403)
    if state['quiz']:
        return fail('Pass this module’s quiz to complete it.', 409)
    completed = parse_bool(request.data.get('completed', True), True)
    m.ChapterProgress.objects.update_or_create(student=request.user, chapter=chapter,
        defaults={'completed': completed, 'completed_at': timezone.now() if completed else None})
    recalc_enrollment(enrollment)
    return Response({'chapterId': chapter.id, 'completed': completed, 'percent': enrollment.percent, 'status': enrollment.status,
                     'currentChapter': enrollment.current_chapter, 'courseCompleted': enrollment.status == 'completed'})


# Resource uploads ------------------------------------------------------------------------------
RESOURCE_EXTENSIONS = {
    'document': {'.pdf', '.doc', '.docx', '.odt', '.txt', '.rtf', '.xls', '.xlsx', '.csv'},
    'presentation': {'.ppt', '.pptx', '.pps', '.ppsx', '.odp', '.pdf'},
}
YOUTUBE = re.compile(r'^https://(www\.|m\.)?(youtube\.com/(watch\?.*v=|embed/|playlist\?.*list=|shorts/)|youtu\.be/)[\w-]+', re.I)


def validate_http_url(url, field='url'):
    try:
        URLValidator(schemes=['http', 'https'])(url)
    except DjangoValidationError:
        raise ValidationError({field: 'Enter a full http(s) link.'})
    return url


def resource_values(data, files, existing=None):
    resource_type = data.get('type', existing.resource_type if existing else 'link')
    if resource_type not in dict(m.CourseContent.RESOURCE_TYPES):
        raise ValidationError({'type': 'Unsupported resource type.'})
    title = str(data.get('title', existing.title if existing else '') or '').strip()
    if not title:
        raise ValidationError({'title': 'Give the resource a title.'})
    url = str(data.get('url', existing.url if existing else '') or '').strip()
    text = str(data.get('text', existing.text if existing else '') or '').strip()[:10000]
    upload = files.get('file')
    has_file = bool(upload) or bool(existing and existing.file)
    if resource_type == 'video':
        if not url or not YOUTUBE.match(url):
            raise ValidationError({'url': 'Use a YouTube video or playlist link (https://www.youtube.com/… or https://youtu.be/…).'})
    elif resource_type == 'link':
        if not url:
            raise ValidationError({'url': 'Enter the link address.'})
        validate_http_url(url)
    elif resource_type in RESOURCE_EXTENSIONS:
        if url:
            validate_http_url(url)
        if not has_file and not url:
            raise ValidationError({'file': 'Upload a file or provide a link to it.'})
        if upload:
            extension = os.path.splitext(upload.name)[1].lower()
            if extension not in RESOURCE_EXTENSIONS[resource_type]:
                raise ValidationError({'file': f'Allowed file types: {", ".join(sorted(RESOURCE_EXTENSIONS[resource_type]))}.'})
            if upload.size > settings.MAX_RESOURCE_UPLOAD_BYTES:
                raise ValidationError({'file': 'Files must be 25 MB or smaller.'})
    elif not text:
        raise ValidationError({'text': 'Write the note text.'})
    if upload and resource_type not in RESOURCE_EXTENSIONS:
        raise ValidationError({'file': 'Files can only be attached to document or presentation resources.'})
    return {'resource_type': resource_type, 'title': title[:200], 'url': url, 'text': text}, upload


@api_view(['GET', 'POST'])
@parser_classes(ANY)
def module_resources(request, pk):
    chapter = get_object_or_404(m.Chapter.objects.select_related('course'), pk=pk)
    if request.method == 'GET':
        if request.user.role == 'student':
            if not get_enrollment(request.user, chapter.course):
                return fail('Enroll in this course to access its resources.', 403)
            if module_state(request.user, chapter)[0]['locked']:
                return fail('Complete the previous required module to unlock this one.', 403)
        return Response([resource_json(item, request.user.role == 'admin') for item in chapter.resources.all()])
    if request.user.role != 'admin':
        return fail('Admin access required.', 403)
    return create_resource(request, chapter.course, chapter)


def create_resource(request, course, chapter=None):
    """Create a module resource, or course-level material when `chapter` is None (same validation and private storage)."""
    values, upload = resource_values(request.data, request.FILES)
    siblings = chapter.resources if chapter else course.contents.filter(chapter__isnull=True)
    item = m.CourseContent(course=course, chapter=chapter, order=siblings.count(), **values)
    if upload:
        item.original_name = get_valid_filename(os.path.basename(upload.name))[:255]
        item.file = upload
    item.save()
    return Response(resource_json(item, True), status=201)


@api_view(['GET', 'POST'])
@parser_classes(ANY)
def course_resources(request, pk):
    """Course-level material (not tied to a module). Students need an enrollment; only administrators add."""
    course = get_object_or_404(m.Course, pk=pk)
    if request.method == 'GET':
        if request.user.role == 'student' and not get_enrollment(request.user, course):
            return fail('Enroll in this course to access its resources.', 403)
        return Response([resource_json(item, request.user.role == 'admin') for item in course.contents.filter(chapter__isnull=True)])
    if request.user.role != 'admin':
        return fail('Admin access required.', 403)
    return create_resource(request, course)


@api_view(['PUT', 'PATCH', 'DELETE'])
@permission_classes([IsAdminRole])
@parser_classes(ANY)
def resource_detail(request, pk):
    item = get_object_or_404(m.CourseContent, pk=pk)
    if request.method == 'DELETE':
        if item.file:
            item.file.delete(save=False)
        item.delete()
        return Response(status=204)
    values, upload = resource_values(request.data, request.FILES, existing=item)
    for field, value in values.items():
        setattr(item, field, value)
    if upload:
        if item.file:
            item.file.delete(save=False)
        item.original_name = get_valid_filename(os.path.basename(upload.name))[:255]
        item.file = upload
    elif parse_bool(request.data.get('removeFile')) and item.file:
        item.file.delete(save=False)
        item.file = None
        item.original_name = ''
    if item.resource_type in RESOURCE_EXTENSIONS and not item.file and not item.url:
        raise ValidationError({'file': 'Upload a file or provide a link to it.'})
    item.save()
    return Response(resource_json(item, True))


@xframe_options_sameorigin
@api_view(['GET'])
def resource_file(request, pk):
    """Stream a private resource file after checking enrollment and module unlock state."""
    item = get_object_or_404(m.CourseContent.objects.select_related('course', 'chapter__course'), pk=pk)
    if not item.file:
        return fail('This resource has no file.', 404)
    if request.user.role == 'student':
        if not get_enrollment(request.user, item.course):
            return fail('Enroll in this course to access its resources.', 403)
        if item.chapter_id and module_state(request.user, item.chapter)[0]['locked']:
            return fail('Complete the previous required module to unlock this resource.', 403)
    name = item.original_name or os.path.basename(item.file.name)
    content_type = mimetypes.guess_type(name)[0] or 'application/octet-stream'
    try:
        handle = item.file.open('rb')
    except Exception as error:  # noqa: BLE001 - never return storage paths to a caller
        logger.error('Course resource storage open failed (%s)', type(error).__name__)
        return fail('The file for this resource is missing on the server.', 404)
    inline = content_type == 'application/pdf' and request.query_params.get('download') != '1'
    response = FileResponse(handle, as_attachment=not inline, filename=name, content_type=content_type)
    response['X-Content-Type-Options'] = 'nosniff'
    response['Cache-Control'] = 'private, no-store'
    response['Content-Security-Policy'] = "default-src 'none'; style-src 'unsafe-inline'; frame-ancestors 'self'"
    return response


# ───────────────────────── Assessment definitions and validation ─────────────────────────
def assessment_type_label(exam):
    return dict(m.Exam.ASSESSMENT_TYPES).get(exam.assessment_type, exam.assessment_type)


def exam_row(exam):
    questions = list(exam.questions.all())
    return {'id': exam.id, 'title': exam.title, 'subject': exam.subject, 'term': exam.term, 'duration': exam.duration,
            'durationSeconds': exam.duration_seconds, 'questions': len(questions), 'points': sum(question.points for question in questions),
            'status': exam.get_status_display(), 'assessmentType': exam.assessment_type, 'assessmentTypeLabel': assessment_type_label(exam),
            'passPercent': int(pass_mark(exam)), 'courseId': exam.course_id, 'courseTitle': exam.course.title if exam.course_id else '',
            'moduleId': exam.module_id, 'moduleTitle': exam.module.title if exam.module else ''}


def parse_choice_reference(raw, choices):
    if isinstance(raw, int):
        return raw
    text = str(raw or '').strip()
    if len(text) == 1 and text.upper() in string.ascii_uppercase:
        return ord(text.upper()) - ord('A')
    try:
        return int(text)
    except ValueError:
        if text in choices:
            return choices.index(text)
    return -1


def split_values(value):
    if isinstance(value, list):
        return value
    if value is None:
        return []
    return [part.strip() for part in re.split(r'[|;,]', str(value)) if part.strip()]


def normalise_answer(value, rules):
    output = str(value or '').strip()
    if rules.get('caseInsensitive', True):
        output = output.casefold()
    if rules.get('ignorePunctuation', False):
        output = ''.join(char for char in output if not unicodedata.category(char).startswith('P'))
    return ' '.join(output.split())


def question_data(raw, index=0):
    question_type = raw.get('questionType', raw.get('question_type', 'single'))
    if question_type not in dict(m.Question.QUESTION_TYPES):
        raise ValidationError({'questions': f'Question {index + 1}: unsupported question type.'})
    text = str(raw.get('text', raw.get('questionText', '')) or '').strip()
    if not text:
        raise ValidationError({'questions': f'Question {index + 1} needs text.'})
    points = to_int(raw.get('points', raw.get('marks', raw.get('score', 1))), 1)
    if points <= 0:
        raise ValidationError({'questions': f'Question {index + 1} has invalid marks.'})
    choices = [str(choice).strip() for choice in (raw.get('choices') or split_values(raw.get('options'))) if str(choice).strip()]
    normalisation = raw.get('normalization') or {}
    if not isinstance(normalisation, dict):
        raise ValidationError({'questions': f'Question {index + 1} has invalid normalization rules.'})
    correct_raw = raw.get('correctAnswers', raw.get('correct', raw.get('correctAnswer')))
    result = {'text': text, 'question_type': question_type, 'choices': choices, 'points': points, 'score': points,
              'order': index, 'normalization': normalisation, 'correct': 0, 'correct_answers': [],
              'expected_answer': '', 'accepted_answers': []}
    if question_type == 'single':
        if len(choices) < 2:
            raise ValidationError({'questions': f'Question {index + 1} needs at least two choices.'})
        refs = split_values(correct_raw)
        if len(refs) != 1:
            raise ValidationError({'questions': f'Question {index + 1} needs exactly one correct choice.'})
        correct = parse_choice_reference(refs[0], choices)
        if correct < 0 or correct >= len(choices):
            raise ValidationError({'questions': f'Question {index + 1} has an invalid correct choice.'})
        result.update(correct=correct, correct_answers=[correct])
    elif question_type == 'multiple':
        if len(choices) < 2:
            raise ValidationError({'questions': f'Question {index + 1} needs at least two choices.'})
        indexes = [parse_choice_reference(value, choices) for value in split_values(correct_raw)]
        if not indexes or any(value < 0 or value >= len(choices) for value in indexes) or len(set(indexes)) != len(indexes):
            raise ValidationError({'questions': f'Question {index + 1} needs valid distinct correct choices.'})
        result['correct_answers'] = sorted(indexes)
    elif question_type == 'true_false':
        if isinstance(correct_raw, list):
            correct_raw = correct_raw[0] if correct_raw else None
        text_value = str(correct_raw).strip().casefold()
        if text_value not in ('true', 'false'):
            raise ValidationError({'questions': f'Question {index + 1} needs a true or false answer.'})
        result.update(choices=['True', 'False'], correct_answers=[text_value == 'true'], correct=0 if text_value == 'true' else 1)
    else:
        expected = str(raw.get('expectedAnswer', raw.get('expected_answer', correct_raw)) or '').strip()
        alternatives = [str(value).strip() for value in split_values(raw.get('acceptedAlternatives', raw.get('accepted_answers'))) if str(value).strip()]
        if not expected:
            raise ValidationError({'questions': f'Question {index + 1} needs an expected short answer.'})
        result.update(expected_answer=expected, accepted_answers=alternatives)
    return result


def check_publishable(exam):
    questions = list(exam.questions.all())
    if not questions:
        raise ValidationError({'message': 'Add at least one question before publishing.'})
    for index, question in enumerate(questions):
        # Re-use authoritative validation so direct API and spreadsheet imports follow one rule.
        question_data({'text': question.text, 'questionType': question.question_type, 'choices': question.choices,
                       'correct': question.correct_answers if question.question_type == 'multiple' else (
                           question.correct_answers[0] if question.correct_answers else question.correct),
                       'expectedAnswer': question.expected_answer, 'acceptedAlternatives': question.accepted_answers,
                       'points': question.points, 'normalization': question.normalization}, index)
    if exam.assessment_type == 'final_exam':
        if not exam.course_id:
            raise ValidationError({'courseId': 'A final exam must belong to a course.'})
        missing = [chapter.title for chapter in exam.course.chapters.filter(required=True)
                   if not m.Exam.objects.filter(course=exam.course, module=chapter, assessment_type='module_quiz', status='published').exists()]
        if missing:
            raise ValidationError({'message': 'Required module quizzes are missing for: ' + ', '.join(missing)})


class DuplicateAssessment(APIException):
    status_code = 409
    default_detail = 'This assessment already exists.'


def save_exam(exam, data, user=None):
    assessment_type = data.get('assessmentType', data.get('assessment_type', exam.assessment_type or 'legacy_quiz'))
    if assessment_type not in dict(m.Exam.ASSESSMENT_TYPES):
        raise ValidationError({'assessmentType': 'Unsupported assessment type.'})
    if data.get('status') and data['status'] not in ('draft', 'published'):
        raise ValidationError({'status': "Must be 'draft' or 'published'."})
    for field in ('title', 'subject', 'term', 'duration', 'status'):
        if field in data:
            setattr(exam, field, data[field])
    exam.assessment_type = assessment_type
    if 'courseId' in data or 'course_id' in data:
        course_id = data.get('courseId', data.get('course_id'))
        exam.course = get_object_or_404(m.Course, pk=course_id) if course_id else None
    if 'moduleId' in data or 'module_id' in data:
        module_id = data.get('moduleId', data.get('module_id'))
        exam.module = get_object_or_404(m.Chapter, pk=module_id) if module_id else None
    if assessment_type == 'final_exam':
        exam.module = None
        exam.duration_seconds = FINAL_DURATION_SECONDS
        exam.duration = '10 mins'
    elif 'durationSeconds' in data:
        exam.duration_seconds = to_int(data.get('durationSeconds'), 0) or None
    if assessment_type == 'final_exam':
        exam.pass_percent = int(PASS_PERCENT)
    elif 'passPercent' in data:
        mark = to_int(data.get('passPercent'), -1)
        if not 1 <= mark <= 100:
            raise ValidationError({'passPercent': 'The pass mark must be between 1 and 100.'})
        exam.pass_percent = mark
    if assessment_type == 'module_quiz':
        if not exam.course_id or not exam.module_id or exam.module.course_id != exam.course_id:
            raise ValidationError({'moduleId': 'A module quiz must belong to a module in its course.'})
    if not exam.title:
        raise ValidationError({'title': 'This field is required.'})
    others = m.Exam.objects.exclude(pk=exam.pk)
    if assessment_type == 'module_quiz' and others.filter(assessment_type='module_quiz', module_id=exam.module_id).exists():
        raise DuplicateAssessment('This module already has a quiz. Edit the existing quiz instead.')
    if assessment_type == 'final_exam' and others.filter(assessment_type='final_exam', course_id=exam.course_id).exists():
        raise DuplicateAssessment('This course already has a final exam. Edit the existing exam instead.')
    if user and not exam.pk:
        exam.created_by = user
    exam.save()
    if 'questions' in data:
        if exam.attempts.filter(status='in_progress').exists():
            raise ValidationError({'message': 'Cannot edit questions while an attempt is in progress.'})
        questions = data['questions'] or []
        validated = [question_data(question, index) for index, question in enumerate(questions)]
        exam.questions.all().delete()
        m.Question.objects.bulk_create([m.Question(exam=exam, **question) for question in validated])
    if exam.status == 'published':
        check_publishable(exam)
    return exam


@api_view(['GET', 'POST'])
def exams(request):
    if request.user.role != 'admin':
        return fail('Admin access required.', 403)
    if request.method == 'GET':
        return Response([exam_row(exam) for exam in m.Exam.objects.select_related('course', 'module').prefetch_related('questions')])
    with transaction.atomic():
        exam = save_exam(m.Exam(), request.data, request.user)
    return Response(exam_row(exam), status=201)


@api_view(['GET', 'PUT', 'DELETE'])
def exam_detail(request, pk):
    if request.user.role != 'admin':
        return fail('Admin access required.', 403)
    exam = get_object_or_404(m.Exam.objects.select_related('course', 'module'), pk=pk)
    if request.method == 'DELETE':
        if exam.attempts.exists():
            return fail('Assessments with attempt history cannot be deleted. Unpublish or archive them instead.', 409)
        with transaction.atomic():
            exam.import_batches.all().delete()
            exam.delete()
        return Response(status=204)
    if request.method == 'PUT':
        with transaction.atomic():
            exam = save_exam(exam, request.data)
    row = exam_row(exam)
    row['attempts'] = exam.attempts.exclude(status='in_progress').count()
    row['questions'] = [{'id': question.id, 'text': question.text, 'questionType': question.question_type,
                         'choices': question.choices, 'correct': question.correct, 'correctAnswers': question.correct_answers,
                         'expectedAnswer': question.expected_answer, 'acceptedAlternatives': question.accepted_answers,
                         'normalization': question.normalization, 'score': question.score, 'points': question.points}
                        for question in exam.questions.all()]
    return Response(row)


# ───────────────────────── Attempt lifecycle, scoring, and finalisation ─────────────────────────
def public_question(question):
    choices = question.choices or (['True', 'False'] if question.question_type == 'true_false' else [])
    return {'id': question.id, 'text': question.text, 'questionType': question.question_type, 'choices': choices, 'points': question.points}


def snapshot_question(question):
    return {**public_question(question), 'correct': question.correct, 'correctAnswers': question.correct_answers,
            'expectedAnswer': question.expected_answer, 'acceptedAlternatives': question.accepted_answers,
            'normalization': question.normalization}


def assessment_payload(exam, student):
    stats = latest_and_best(student, exam)
    return {'id': exam.id, 'title': exam.title, 'subject': exam.subject, 'assessmentType': exam.assessment_type,
            'durationSeconds': FINAL_DURATION_SECONDS if exam.assessment_type == 'final_exam' else exam.duration_seconds,
            'duration': '10 mins' if exam.assessment_type == 'final_exam' else exam.duration,
            'passPercent': int(pass_mark(exam)), 'courseTitle': exam.course.title if exam.course_id else '',
            'moduleTitle': exam.module.title if exam.module_id else '',
            'questionCount': exam.questions.count(), 'totalPoints': sum(question.points for question in exam.questions.all()),
            'courseId': exam.course_id, 'moduleId': exam.module_id, **stats,
            'eligibility': final_eligibility(student, exam)}


def student_visible_assessments(student):
    exams = m.Exam.objects.filter(status='published').exclude(assessment_type='final_exam').select_related('course', 'module').prefetch_related('questions')
    visible = []
    for exam in exams:
        if exam.course_id and not get_enrollment(student, exam.course):
            continue
        visible.append(exam)
    return visible


def attempt_response(attempt, include_correct=False):
    # serverTime lets the browser correct for clock skew; the deadline itself is enforced here, not in React.
    payload = {'attemptId': attempt.id, 'assessmentId': attempt.exam_id, 'status': attempt.status,
               'startedAt': attempt.started_at, 'deadline': attempt.deadline_at, 'submittedAt': attempt.submitted_at,
               'serverTime': timezone.now(), 'passPercent': int(pass_mark(attempt.exam)),
               'answers': attempt.answers, 'answerVersion': attempt.answer_version, 'score': value_number(attempt.score),
               'autoScore': value_number(attempt.auto_score), 'total': value_number(attempt.total_points),
               'percent': value_number(attempt.percentage), 'passed': attempt.passed,
               'questions': [{key: value for key, value in question.items() if include_correct or key not in ('correct', 'correctAnswers', 'expectedAnswer', 'acceptedAlternatives', 'normalization')}
                             for question in attempt.question_snapshot]}
    payload['questionCount'] = len(attempt.question_snapshot)
    if include_correct and attempt.status != 'in_progress':
        payload['review'] = grade_snapshot(attempt.question_snapshot, attempt.answers)[2]
        # Questions can carry different points, so the percentage is points earned / points possible;
        # the correct-question count is reported separately so the two are never confused.
        payload['correctCount'] = sum(1 for item in payload['review'] if item['isCorrect'])
    return payload


def grade_snapshot(questions, answers):
    total = Decimal('0')
    score = Decimal('0')
    review = []
    for question in questions:
        points = Decimal(str(question.get('points', 0)))
        total += points
        question_id = str(question['id'])
        chosen = answers.get(question_id)
        question_type = question.get('questionType', 'single')
        if question_type == 'single':
            correct = (question.get('correctAnswers') or [question.get('correct', 0)])[0]
            ok = to_int(chosen, -1) == int(correct)
        elif question_type == 'multiple':
            expected = sorted({to_int(value, -1) for value in question.get('correctAnswers', [])})
            actual = sorted({to_int(value, -1) for value in (chosen if isinstance(chosen, list) else [])})
            ok = actual == expected
        elif question_type == 'true_false':
            expected = bool((question.get('correctAnswers') or [False])[0])
            actual = chosen if isinstance(chosen, bool) else str(chosen).strip().casefold() == 'true'
            ok = actual == expected
        else:
            rules = question.get('normalization') or {}
            candidates = [question.get('expectedAnswer', '')] + (question.get('acceptedAlternatives') or [])
            ok = any(normalise_answer(chosen, rules) == normalise_answer(candidate, rules) for candidate in candidates)
        earned = points if ok else Decimal('0')
        score += earned
        review.append({'id': question['id'], 'text': question['text'], 'questionType': question_type,
                       'choices': question.get('choices', []), 'chosen': chosen, 'correct': question.get('correct'),
                       'correctAnswers': question.get('correctAnswers', []), 'expectedAnswer': question.get('expectedAnswer', ''),
                       'isCorrect': ok, 'points': value_number(points), 'earned': value_number(earned)})
    return (score, total, review)


def update_completion(attempt):
    """Record earned progress. A later failed retake deliberately cannot undo an earlier pass."""
    exam = attempt.exam
    if not exam.course_id or not attempt.passed:
        return
    enrollment = m.Enrollment.objects.filter(student=attempt.student, course=exam.course).select_related('course', 'student').first()
    if exam.assessment_type == 'module_quiz':
        if exam.module_id:
            m.ChapterProgress.objects.update_or_create(student=attempt.student, chapter_id=exam.module_id,
                                                       defaults={'completed': True, 'completed_at': attempt.submitted_at or timezone.now()})
        if enrollment:
            recalc_enrollment(enrollment)
        return
    if exam.assessment_type != 'final_exam':
        return
    if enrollment and enrollment.status != 'completed':
        enrollment.status = 'completed'
        enrollment.completed_at = timezone.now()
        enrollment.completion_source = 'final_exam'
        enrollment.save(update_fields=['status', 'completed_at', 'completion_source'])


def finalize_attempt(attempt_id, expired=False):
    """Idempotent finaliser shared by expiry processing and manual submit."""
    with transaction.atomic():
        attempt = m.Attempt.objects.select_for_update(of=('self',)).select_related('exam', 'student', 'exam__course').get(pk=attempt_id)
        if attempt.status != 'in_progress':
            return attempt
        now = timezone.now()
        is_expired = expired or bool(attempt.deadline_at and attempt.deadline_at <= now)
        score, total, _ = grade_snapshot(attempt.question_snapshot, attempt.answers)
        percent = (score * Decimal('100') / total).quantize(Decimal('0.01')) if total else Decimal('0')
        attempt.auto_score = score
        attempt.score = score
        attempt.total_points = total
        attempt.percentage = percent
        attempt.passed = percent >= pass_mark(attempt.exam)
        attempt.status = 'expired' if is_expired else 'submitted'
        attempt.submitted_at = now
        attempt.save(update_fields=['auto_score', 'score', 'total_points', 'percentage', 'passed', 'status', 'submitted_at'])
        update_completion(attempt)
        return attempt


def finalize_expired_attempts():
    """Called at request boundaries and by the scheduled management command."""
    ids = list(m.Attempt.objects.filter(status='in_progress', deadline_at__lte=timezone.now()).values_list('id', flat=True))
    return [finalize_attempt(attempt_id, expired=True).id for attempt_id in ids]


def create_or_resume_attempt(student, exam):
    if student.role != 'student':
        raise ValidationError({'message': 'Only students can take assessments.'})
    if exam.status != 'published':
        raise ValidationError({'message': 'This assessment is not published.'})
    if exam.course_id and not get_enrollment(student, exam.course):
        raise ValidationError({'message': 'Enroll in this course before taking its assessment.'})
    eligibility = final_eligibility(student, exam)
    if not eligibility['eligible']:
        raise ValidationError({'message': ' '.join(eligibility['reasons'])})
    if exam.assessment_type == 'module_quiz' and exam.module_id and module_state(student, exam.module)[0]['locked']:
        raise ValidationError({'message': 'Complete the previous required module to unlock this quiz.'})
    finalize_expired_attempts()
    with transaction.atomic():
        existing = m.Attempt.objects.select_for_update().filter(student=student, exam=exam, status='in_progress').first()
        if existing:
            return existing, False
        now = timezone.now()
        snapshot = [snapshot_question(question) for question in exam.questions.all()]
        if not snapshot:
            raise ValidationError({'message': 'This assessment has no questions.'})
        try:
            # The insert gets its own savepoint: on PostgreSQL a failed statement aborts the enclosing
            # transaction, so without it the lookup below raised TransactionManagementError (HTTP 500).
            with transaction.atomic():
                attempt = m.Attempt.objects.create(student=student, exam=exam, started_at=now,
                    deadline_at=now + timedelta(seconds=FINAL_DURATION_SECONDS) if exam.assessment_type == 'final_exam' else None,
                    question_snapshot=snapshot, course_snapshot={'id': exam.course_id, 'title': exam.course.title if exam.course_id else ''})
        except IntegrityError:
            # A concurrent start (double click, React StrictMode remount, second tab) committed first:
            # the partial unique constraint guarantees one active attempt, so resume that one.
            attempt = m.Attempt.objects.select_related('exam').get(student=student, exam=exam, status='in_progress')
            return attempt, False
        return attempt, True


def get_owned_attempt(request, pk):
    attempt = get_object_or_404(m.Attempt.objects.select_related('exam', 'exam__course'), pk=pk)
    if request.user.role == 'student' and attempt.student_id != request.user.id:
        return None
    return attempt


@api_view(['GET'])
@permission_classes([IsStudentRole])
def assessments(request):
    if request.user.role != 'student':
        return fail('Only students can view assessments.', 403)
    finalize_expired_attempts()
    items = []
    for exam in student_visible_assessments(request.user):
        items.append(assessment_payload(exam, request.user))
    for exam in m.Exam.objects.filter(status='published', assessment_type='final_exam').select_related('course').prefetch_related('questions'):
        if exam.course_id and get_enrollment(request.user, exam.course):
            items.append(assessment_payload(exam, request.user))
    return Response(items)


def student_assessment_detail(request, pk):
    if request.user.role != 'student':
        return fail('Only students can view assessments.', 403)
    exam = get_object_or_404(m.Exam.objects.prefetch_related('questions'), pk=pk, status='published')
    if exam.course_id and not get_enrollment(request.user, exam.course):
        return fail('Enroll in this course before viewing this assessment.', 403)
    eligibility = final_eligibility(request.user, exam)
    if exam.assessment_type == 'final_exam' and not eligibility['eligible']:
        return fail(' '.join(eligibility['reasons']), 403)
    if exam.assessment_type == 'module_quiz' and exam.module_id and module_state(request.user, exam.module)[0]['locked']:
        return fail('Complete the previous required module to unlock this quiz.', 403)
    finalize_expired_attempts()
    active = m.Attempt.objects.filter(student=request.user, exam=exam, status='in_progress').first()
    return Response({**assessment_payload(exam, request.user), 'activeAttemptId': active.id if active else None,
                     'activeDeadline': iso(active.deadline_at) if active else None,
                     'questions': [public_question(question) for question in exam.questions.all()]})


@api_view(['GET'])
@permission_classes([IsStudentRole])
def assessment_detail(request, pk):
    return student_assessment_detail(request, pk)


@api_view(['POST'])
@permission_classes([IsStudentRole])
def assessment_start(request, pk):
    exam = get_object_or_404(m.Exam.objects.prefetch_related('questions').select_related('course'), pk=pk)
    try:
        attempt, created = create_or_resume_attempt(request.user, exam)
    except ValidationError as error:
        return fail(str(error.detail.get('message', error.detail)), 403)
    response = attempt_response(attempt)
    response['resumed'] = not created
    return Response(response, status=201 if created else 200)


@api_view(['GET'])
@permission_classes([IsStudentRole])
def attempt_detail(request, pk):
    finalize_expired_attempts()
    attempt = get_owned_attempt(request, pk)
    if not attempt:
        return fail('You cannot access this attempt.', 403)
    # Correct answers are only revealed once the attempt has been finalized.
    return Response(attempt_response(attempt, include_correct=attempt.status != 'in_progress'))


@api_view(['PUT'])
@permission_classes([IsStudentRole])
def attempt_answers(request, pk):
    finalize_expired_attempts()
    attempt = get_owned_attempt(request, pk)
    if not attempt:
        return fail('You cannot modify this attempt.', 403)
    if attempt.status != 'in_progress':
        return fail('This attempt has already been finalized.', 409)
    version = request.data.get('version')
    if version is None or to_int(version, -1) != attempt.answer_version:
        return Response({'message': 'Answers changed elsewhere. Refresh the attempt before saving.', 'answerVersion': attempt.answer_version,
                         'answers': attempt.answers}, status=409)
    updates = request.data.get('answers')
    if not isinstance(updates, dict):
        raise ValidationError({'answers': 'Must be an object keyed by question id.'})
    valid_ids = {str(question['id']) for question in attempt.question_snapshot}
    if not set(str(key) for key in updates).issubset(valid_ids):
        raise ValidationError({'answers': 'An answer references a question outside this attempt.'})
    merged = dict(attempt.answers)
    merged.update({str(key): value for key, value in updates.items()})
    attempt.answers = merged
    attempt.answer_version += 1
    attempt.save(update_fields=['answers', 'answer_version'])
    return Response({'answers': attempt.answers, 'answerVersion': attempt.answer_version, 'savedAt': timezone.now(),
                     'serverTime': timezone.now(), 'deadline': attempt.deadline_at})


@api_view(['POST'])
@permission_classes([IsStudentRole])
def attempt_submit(request, pk):
    attempt = get_owned_attempt(request, pk)
    if not attempt:
        return fail('You cannot submit this attempt.', 403)
    final = finalize_attempt(attempt.id)
    return Response(attempt_response(final, include_correct=True))


@api_view(['GET'])
@permission_classes([IsStudentRole])
def assessment_history(request, pk):
    exam = get_object_or_404(m.Exam, pk=pk)
    if request.user.role != 'student' or (exam.course_id and not get_enrollment(request.user, exam.course)):
        return fail('You cannot access this assessment history.', 403)
    attempts = m.Attempt.objects.filter(student=request.user, exam=exam).exclude(status='in_progress').order_by('-submitted_at', '-id')
    return Response([{'id': attempt.id, 'status': attempt.status, 'score': value_number(attempt.score), 'total': value_number(attempt.total_points),
                      'percent': value_number(attempt.percentage), 'passed': attempt.passed, 'submittedAt': attempt.submitted_at,
                      'startedAt': attempt.started_at} for attempt in attempts])


# Compatibility API: legacy course screens keep working while now using server-side lifecycle rules.
@api_view(['GET'])
@permission_classes([IsStudentRole])
def quizzes(request):
    if request.user.role != 'student':
        return fail('Only students can view quizzes.', 403)
    return Response([{**assessment_payload(exam, request.user), 'questions': exam.questions.count(),
                      'points': sum(question.points for question in exam.questions.all())}
                     for exam in student_visible_assessments(request.user)])


@api_view(['GET'])
@permission_classes([IsStudentRole])
def quiz_detail(request, pk):
    return student_assessment_detail(request, pk)


@api_view(['POST'])
@permission_classes([IsStudentRole])
def quiz_attempt(request, pk):
    exam = get_object_or_404(m.Exam.objects.prefetch_related('questions').select_related('course'), pk=pk)
    if exam.assessment_type == 'final_exam':
        return fail('Use the final-exam start endpoint for timed final examinations.', 400)
    try:
        attempt, _ = create_or_resume_attempt(request.user, exam)
    except ValidationError as error:
        return fail(str(error.detail.get('message', error.detail)), 403)
    answers = request.data.get('answers')
    if not isinstance(answers, dict):
        raise ValidationError({'answers': 'Must be an object keyed by question id.'})
    valid_ids = {str(question['id']) for question in attempt.question_snapshot}
    if not {str(key) for key in answers} <= valid_ids:
        raise ValidationError({'answers': 'An answer references a question outside this attempt.'})
    attempt.answers = {str(key): value for key, value in answers.items()}
    attempt.answer_version += 1
    attempt.save(update_fields=['answers', 'answer_version'])
    final = finalize_attempt(attempt.id)
    response = attempt_response(final, include_correct=True)
    stats = latest_and_best(request.user, exam)
    # `passed` stays this attempt's own outcome; `everPassed` reports whether any attempt has passed.
    response.update({key: value for key, value in stats.items() if key != 'passed'}, everPassed=stats['passed'])
    return Response(response, status=201)


# ───────────────────────── Student directory and analytics ─────────────────────────
AUTH_METHOD_LABELS = dict(m.LoginEvent.AUTH_METHODS)
STUDENT_SORTS = {'name': ('first_name', 'last_name', 'username'), 'registered': ('date_joined',), 'lastLogin': ('last_login_at',),
                 'courses': ('enrollment_count',), 'progress': ('average_progress',)}


def directory_queryset():
    """Students annotated with summary numbers through correlated subqueries (one query, no join fan-out)."""
    enrollments = m.Enrollment.objects.filter(student=OuterRef('pk')).values('student')
    logins = m.LoginEvent.objects.filter(user=OuterRef('pk'), event_type='login')
    return m.User.objects.filter(role='student').annotate(
        enrollment_count=Coalesce(Subquery(enrollments.annotate(c=Count('id')).values('c')[:1], output_field=IntegerField()), Value(0)),
        completed_count=Coalesce(Subquery(enrollments.filter(status='completed').annotate(c=Count('id')).values('c')[:1],
                                          output_field=IntegerField()), Value(0)),
        average_progress=Subquery(enrollments.annotate(a=Avg('percent')).values('a')[:1]),
        login_count=Coalesce(Subquery(logins.values('user').annotate(c=Count('id')).values('c')[:1], output_field=IntegerField()), Value(0)),
        last_login_at=Subquery(logins.order_by('-created_at').values('created_at')[:1]),
    )


def directory_row(user):
    return {'id': user.id, 'name': user.display_name, 'username': user.username, 'email': user.email, 'isActive': user.is_active,
            'registeredAt': iso(user.date_joined), 'lastLoginAt': iso(user.last_login_at), 'loginCount': user.login_count,
            'enrollmentCount': user.enrollment_count, 'completedCount': user.completed_count,
            'averageProgress': round(user.average_progress) if user.average_progress is not None else None,
            'googleLinked': bool(user.google_sub)}


def attempt_history_row(attempt):
    return {'id': attempt.id, 'assessmentId': attempt.exam_id, 'title': attempt.exam.title, 'type': attempt.exam.assessment_type,
            'courseId': attempt.exam.course_id,
            'course': attempt.course_snapshot.get('title') or (attempt.exam.course.title if attempt.exam.course_id else ''),
            'moduleTitle': attempt.exam.module.title if attempt.exam.module_id else '', 'status': attempt.status,
            'score': value_number(attempt.score), 'total': value_number(attempt.total_points), 'percent': value_number(attempt.percentage),
            'passed': attempt.passed, 'startedAt': iso(attempt.started_at), 'submittedAt': iso(attempt.submitted_at),
            'reviewed': bool(attempt.grading_reviews.all())}


def student_record(user):
    """Admin view of one student. Never includes password hashes, sessions, tokens or OAuth subjects."""
    row = directory_row(directory_queryset().get(pk=user.pk))
    enrollments = [student_course_report(user, enrollment) for enrollment in user.enrollments.select_related('course').order_by('-id')]
    attempts = list(m.Attempt.objects.filter(student=user).select_related('exam', 'exam__course', 'exam__module')
                    .prefetch_related('grading_reviews').order_by('-created_at'))
    finalized = [attempt for attempt in attempts if attempt.status in FINALIZED]
    module_attempts = [attempt for attempt in finalized if attempt.exam.assessment_type != 'final_exam']
    finals = [attempt for attempt in finalized if attempt.exam.assessment_type == 'final_exam']
    logins = user.login_events.filter(event_type='login').order_by('-created_at')[:25]
    return {**row, 'firstName': user.first_name, 'lastName': user.last_name, 'city': user.city, 'country': user.country,
            'enrollments': enrollments,
            'metrics': {'quizAttempts': len(module_attempts), 'quizzesPassed': len({a.exam_id for a in module_attempts if a.passed}),
                        'averageQuizPercent': (round(sum(float(a.percentage) for a in module_attempts) / len(module_attempts))
                                               if module_attempts else None),
                        'finalAttempts': len(finals), 'finalsPassed': len({a.exam_id for a in finals if a.passed}),
                        'inProgressAttempts': sum(1 for a in attempts if a.status == 'in_progress')},
            'attempts': [attempt_history_row(attempt) for attempt in finalized],
            'loginHistory': [{'id': event.id, 'createdAt': iso(event.created_at), 'authMethod': event.auth_method,
                              'authMethodLabel': AUTH_METHOD_LABELS.get(event.auth_method, 'Unspecified')} for event in logins],
            'lastLogin': iso(user.last_login), 'loginEvents': row['loginCount']}


@api_view(['GET'])
@permission_classes([IsAdminRole])
def students(request):
    params = request.query_params
    queryset = directory_queryset()
    query = (params.get('q') or '').strip()
    if query:
        queryset = queryset.filter(Q(username__icontains=query) | Q(email__icontains=query) |
                                   Q(first_name__icontains=query) | Q(last_name__icontains=query))
    status_value = params.get('status')
    if status_value in ('ongoing', 'completed'):
        queryset = queryset.filter(pk__in=m.Enrollment.objects.filter(status=status_value).values('student_id'))
    enrollment_value = params.get('enrollment')
    if enrollment_value == 'enrolled':
        queryset = queryset.filter(enrollment_count__gt=0)
    elif enrollment_value == 'not_enrolled':
        queryset = queryset.filter(enrollment_count=0)
    login_value = params.get('login')
    if login_value == 'logged_in':
        queryset = queryset.filter(login_count__gt=0)
    elif login_value == 'never':
        queryset = queryset.filter(login_count=0)
    sort = params.get('sort') or 'name'
    descending = sort.startswith('-')
    fields = STUDENT_SORTS.get(sort.lstrip('-'))
    if not fields:
        raise ValidationError({'sort': 'Unsupported sort.'})
    ordering = [(F(field).desc(nulls_last=True) if descending else F(field).asc(nulls_last=True)) for field in fields]
    queryset = queryset.order_by(*ordering, 'id')
    page = max(to_int(params.get('page'), 1), 1)
    page_size = min(max(to_int(params.get('pageSize'), 20), 1), 100)
    total = queryset.count()
    pages = max((total + page_size - 1) // page_size, 1)
    page = min(page, pages)
    rows = queryset[(page - 1) * page_size: page * page_size]
    all_students = m.User.objects.filter(role='student')
    week_ago = timezone.now() - timedelta(days=7)
    summary = {'total': all_students.count(), 'enrolled': all_students.filter(enrollments__isnull=False).distinct().count(),
               'neverLoggedIn': all_students.exclude(login_events__event_type='login').count(),
               'activeLast7Days': m.LoginEvent.objects.filter(event_type='login', user__role='student', created_at__gte=week_ago)
               .values('user_id').distinct().count()}
    return Response({'results': [directory_row(user) for user in rows], 'page': page, 'pageSize': page_size, 'total': total,
                     'pages': pages, 'summary': summary})


@api_view(['GET'])
@permission_classes([IsAdminRole])
def student_detail(request, pk):
    student = get_object_or_404(m.User, pk=pk, role='student')
    finalize_expired_attempts()
    return Response(student_record(student))


@api_view(['GET'])
@permission_classes([IsAdminRole])
def admin_attempt_detail(request, pk):
    attempt = get_object_or_404(m.Attempt.objects.select_related('exam', 'student'), pk=pk)
    payload = attempt_response(attempt, include_correct=attempt.status != 'in_progress')
    payload.update({'student': {'id': attempt.student_id, 'name': attempt.student.display_name}, 'title': attempt.exam.title,
                    'reviews': [{'id': review.id, 'originalScore': value_number(review.original_score),
                                 'revisedScore': value_number(review.revised_score), 'reason': review.reason,
                                 'reviewer': review.reviewer.display_name, 'createdAt': iso(review.created_at)}
                                for review in attempt.grading_reviews.select_related('reviewer')]})
    return Response(payload)


@api_view(['GET'])
@permission_classes([IsAdminRole])
def login_analytics(request):
    events = date_filter(request, m.LoginEvent.objects.select_related('user').filter(event_type='login'))
    page = max(to_int(request.query_params.get('page'), 1), 1)
    page_size = min(max(to_int(request.query_params.get('pageSize'), 50), 1), 100)
    total = events.count()
    rows = events[(page - 1) * page_size: page * page_size]
    return Response({'metrics': {'registeredStudents': m.User.objects.filter(role='student').count(),
        'newRegistrations': date_filter(request, m.User.objects.filter(role='student'), 'date_joined').count(),
        'uniqueSignedIn': events.filter(event_type='login').values('user_id').distinct().count(),
        'successfulLogins': events.count(),
        'emailPasswordLogins': events.filter(auth_method='email_password').count(),
        'googleLogins': events.filter(auth_method='google').count(),
        'legacyJwtLogins': events.filter(auth_method='legacy_jwt').count()},
        'events': [{'id': event.id, 'userId': event.user_id, 'name': event.user.display_name, 'eventType': event.event_type,
                    'authMethod': event.auth_method, 'createdAt': event.created_at} for event in rows], 'page': page, 'pageSize': page_size, 'total': total})


@api_view(['GET'])
@permission_classes([IsAdminRole])
def exam_analytics(request):
    attempts = date_filter(request, m.Attempt.objects.select_related('exam').filter(exam__assessment_type='final_exam'))
    completed = attempts.exclude(status='in_progress')
    published = m.Exam.objects.filter(assessment_type='final_exam', status='published').count()
    passed = completed.filter(passed=True).count()
    total_completed = completed.count()
    average = completed.aggregate(value=Avg('percentage'))['value'] or Decimal('0')
    return Response({'finalExamDurationSeconds': FINAL_DURATION_SECONDS, 'metrics': {'publishedFinalExams': published,
        'totalAttempts': attempts.count(), 'inProgress': attempts.filter(status='in_progress').count(), 'completed': total_completed,
        'expired': attempts.filter(status='expired').count(), 'passed': passed, 'failed': total_completed - passed,
        'passRate': float(Decimal('100') * passed / total_completed) if total_completed else 0, 'averagePercent': value_number(average)}})


@api_view(['POST'])
@permission_classes([IsAdminRole])
def review_attempt(request, pk):
    attempt = get_object_or_404(m.Attempt.objects.select_related('exam', 'student'), pk=pk)
    if attempt.status == 'in_progress':
        return fail('An in-progress attempt cannot be reviewed.', 409)
    reason = (request.data.get('reason') or '').strip()
    if not reason:
        raise ValidationError({'reason': 'A review reason is required.'})
    revised = request.data.get('revisedScore')
    try:
        revised_score = Decimal(str(revised))
    except Exception:
        raise ValidationError({'revisedScore': 'Use a numeric score.'})
    if revised_score < 0 or revised_score > attempt.total_points:
        raise ValidationError({'revisedScore': 'Score must be between zero and the assessment total.'})
    with transaction.atomic():
        original = attempt.score
        attempt.score = revised_score
        attempt.percentage = (revised_score * Decimal('100') / attempt.total_points).quantize(Decimal('0.01')) if attempt.total_points else 0
        attempt.passed = attempt.percentage >= pass_mark(attempt.exam)
        attempt.save(update_fields=['score', 'percentage', 'passed'])
        review = m.GradingReview.objects.create(attempt=attempt, reviewer=request.user, original_score=original,
                                                 revised_score=revised_score, reason=reason,
                                                 adjustments=request.data.get('adjustments') or {})
        update_completion(attempt)
    return Response({'reviewId': review.id, 'attempt': attempt_response(attempt, include_correct=True)})


# ───────────────────────── Excel question import ─────────────────────────
def get_openpyxl():
    try:
        from openpyxl import load_workbook
        return load_workbook
    except ImportError:
        raise ValidationError({'file': 'Excel imports require the openpyxl dependency.'})


def serialise_cell(value):
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return value


@api_view(['POST'])
@permission_classes([IsAdminRole])
@parser_classes([MultiPartParser])
def question_import_upload(request):
    upload = request.FILES.get('file')
    if not upload:
        raise ValidationError({'file': 'Select an .xlsx spreadsheet.'})
    if not upload.name.lower().endswith('.xlsx') or upload.size > MAX_IMPORT_BYTES:
        raise ValidationError({'file': 'Use an .xlsx file smaller than 5 MB.'})
    content = upload.read()
    # .xlsx files are ZIP packages; anything else (renamed .xls/.csv, HTML exports) is rejected before parsing.
    if not content.startswith(b'PK\x03\x04'):
        raise ValidationError({'file': 'This file is not a valid .xlsx workbook. Save it from Excel as “Excel Workbook (.xlsx)”.'})
    exam = get_object_or_404(m.Exam, pk=to_int(request.data.get('examId'), -1))
    course_id = request.data.get('courseId')
    if course_id and str(exam.course_id or '') != str(course_id):
        raise ValidationError({'courseId': 'The selected assessment does not belong to this course.'})
    try:
        workbook = get_openpyxl()(io.BytesIO(content), read_only=True, data_only=True, keep_vba=False)
        sheets = {}
        for sheet in workbook.worksheets:
            rows = [[serialise_cell(cell) for cell in row] for row in sheet.iter_rows(values_only=True)]
            headers = [str(value).strip() if value is not None else '' for value in (rows[0] if rows else [])]
            sheets[sheet.title] = {'headers': headers, 'rows': rows[1:] if rows else []}
    except Exception:
        raise ValidationError({'file': 'The workbook could not be read as a safe .xlsx file.'})
    if not any(any(data['headers']) and data['rows'] for data in sheets.values()):
        raise ValidationError({'file': 'The workbook has no worksheet with a header row and at least one question row.'})
    batch = m.QuestionImportBatch.objects.create(public_id=uuid.uuid4().hex, admin=request.user, exam=exam, workbook={'sheets': sheets})
    return Response({'batchId': batch.public_id, 'exam': {'id': exam.id, 'title': exam.title, 'questionCount': exam.questions.count()},
                     'sheets': [{'name': name, 'headers': data['headers'], 'rows': len(data['rows'])} for name, data in sheets.items()]}, status=201)


IMPORT_FIELDS = ('questionType', 'questionText', 'options', 'correctAnswer', 'marks', 'acceptedAlternatives')


def public_candidate(candidate):
    """Preview rows for the administrator, in the same shape the question editor uses."""
    return {'text': candidate['text'], 'questionType': candidate['question_type'], 'choices': candidate['choices'],
            'points': candidate['points'], 'correctAnswers': candidate['correct_answers'],
            'expectedAnswer': candidate['expected_answer'], 'acceptedAlternatives': candidate['accepted_answers']}


def mapped_row(headers, row, mapping):
    values = {headers[index]: row[index] if index < len(row) else None for index in range(len(headers))}
    return {field: values.get(column) for field, column in mapping.items()}


@api_view(['POST'])
@permission_classes([IsAdminRole])
def question_import_preview(request, batch_id):
    batch = get_object_or_404(m.QuestionImportBatch, public_id=batch_id, admin=request.user)
    sheet_name = request.data.get('sheet')
    mapping = request.data.get('mapping')
    sheets = batch.workbook.get('sheets', {})
    if sheet_name not in sheets or not isinstance(mapping, dict):
        raise ValidationError({'message': 'Choose a worksheet and column mapping.'})
    sheet = sheets[sheet_name]
    headers = sheet['headers']
    errors, candidates, seen = [], [], set()
    required_map = {'questionType', 'questionText', 'marks'}
    missing = [field for field in sorted(required_map) if not mapping.get(field)]
    if missing:
        raise ValidationError({'mapping': 'Missing mappings: ' + ', '.join(missing)})
    unknown = sorted(str(column) for field, column in mapping.items() if column and column not in headers)
    if unknown:
        raise ValidationError({'mapping': 'These columns are not in the worksheet: ' + ', '.join(unknown)})
    mapping = {field: column for field, column in mapping.items() if field in IMPORT_FIELDS and column}
    blank_rows = []
    for row_number, row in enumerate(sheet['rows'], start=2):
        if not any(value not in (None, '') for value in row):
            # Completely empty rows (often left by Excel formatting) hold no question; they are reported, not imported.
            blank_rows.append(row_number)
            continue
        source = mapped_row(headers, row, mapping)
        raw = {'questionType': source.get('questionType'), 'text': source.get('questionText'), 'options': source.get('options'),
               'correctAnswer': source.get('correctAnswer'), 'marks': source.get('marks'),
               'acceptedAlternatives': source.get('acceptedAlternatives'), 'normalization': {}}
        try:
            candidate = question_data(raw, row_number - 2)
            signature = json.dumps({'text': candidate['text'].casefold(), 'type': candidate['question_type'],
                                    'choices': candidate['choices']}, sort_keys=True)
            if signature in seen or m.Question.objects.filter(exam=batch.exam, text__iexact=candidate['text']).exists():
                raise ValidationError({'questions': f'Question {row_number}: duplicate candidate.'})
            seen.add(signature)
            candidates.append(candidate)
        except ValidationError as error:
            errors.append({'row': row_number, 'message': str(error.detail.get('questions', error.detail))})
    if not candidates and not errors:
        raise ValidationError({'message': 'This worksheet has no question rows.'})
    validation = {'sheet': sheet_name, 'headers': headers, 'candidates': candidates, 'errors': errors}
    batch.mapping, batch.validation, batch.status = mapping, validation, 'previewed'
    batch.save(update_fields=['mapping', 'validation', 'status'])
    return Response({'validCount': len(candidates), 'invalidCount': len(errors), 'preview': [public_candidate(item) for item in candidates],
                     'errors': errors, 'blankRows': blank_rows})


@api_view(['POST'])
@permission_classes([IsAdminRole])
def question_import_confirm(request, batch_id):
    with transaction.atomic():
        batch = get_object_or_404(m.QuestionImportBatch.objects.select_for_update(), public_id=batch_id, admin=request.user)
        key = (request.data.get('idempotencyKey') or '').strip()
        if not key:
            raise ValidationError({'idempotencyKey': 'An idempotency key is required.'})
        if batch.status == 'imported':
            if key == batch.idempotency_key:
                return Response({'imported': len(batch.validation.get('candidates', [])), 'rejected': len(batch.validation.get('errors', [])), 'alreadyImported': True})
            return fail('This import batch was already confirmed.', 409)
        candidates, errors = batch.validation.get('candidates', []), batch.validation.get('errors', [])
        if batch.status != 'previewed' or errors:
            return fail('Resolve validation errors and preview again before importing.', 400)
        exam = m.Exam.objects.select_for_update().get(pk=batch.exam_id)
        if exam.attempts.filter(status='in_progress').exists():
            return fail('A student is taking this assessment right now. Import the questions after the attempt ends.', 409)
        # Another batch may have added the same questions since this preview; never import them twice.
        existing = {text.casefold() for text in exam.questions.values_list('text', flat=True)}
        if any(candidate['text'].casefold() in existing for candidate in candidates):
            return fail('Some of these questions were added to the assessment after the preview. Preview again before importing.', 409)
        highest = exam.questions.aggregate(value=Max('order'))['value']
        start = 0 if highest is None else highest + 1
        m.Question.objects.bulk_create([m.Question(exam=exam, **{**candidate, 'order': start + index})
                                        for index, candidate in enumerate(candidates)])
        batch.status, batch.idempotency_key, batch.confirmed_at = 'imported', key, timezone.now()
        batch.save(update_fields=['status', 'idempotency_key', 'confirmed_at'])
    return Response({'imported': len(candidates), 'rejected': 0, 'alreadyImported': False}, status=201)


# ───────────────────────── Library, notes, calendar and profile ─────────────────────────
@api_view(['GET'])
def library(request):
    return Response([{'id': item.id, 'title': item.title, 'subject': item.subject, 'cover': abs_url(request, item.cover),
                      'fileUrl': abs_url(request, item.file)} for item in m.LibraryItem.objects.all()])


NOTES_APPS = [
    {'id': 'drive', 'name': 'Google Drive', 'url': 'https://drive.google.com', 'color': '#34a853'},
    {'id': 'lens', 'name': 'Google Lens', 'url': 'https://lens.google.com', 'color': '#4285f4'},
    {'id': 'excel', 'name': 'Excel', 'url': 'https://www.office.com/launch/excel', 'color': '#107c41'},
    {'id': 'doc', 'name': 'Google Doc', 'url': 'https://docs.google.com', 'color': '#4285f4'},
    {'id': 'word', 'name': 'Word', 'url': 'https://www.office.com/launch/word', 'color': '#2b579a'},
    {'id': 'ppt', 'name': 'Power point', 'url': 'https://www.office.com/launch/powerpoint', 'color': '#d24726'},
]


@api_view(['GET'])
def notes_apps(request):
    return Response(NOTES_APPS)


@api_view(['GET', 'POST'])
def calendar_view(request):
    if request.method == 'POST':
        data = request.data
        if not data.get('date') or not data.get('title'):
            raise ValidationError({'message': 'date and title are required.'})
        event = m.CalendarEvent.objects.create(user=request.user, date=data['date'], title=data['title'], time=data.get('time', ''), type=data.get('type', 'event'))
        return Response({'id': event.id, 'date': str(event.date), 'title': event.title, 'type': event.type}, status=201)
    today = date.today()
    year, month = to_int(request.query_params.get('year'), today.year), to_int(request.query_params.get('month'), today.month)
    if not 1 <= month <= 12:
        raise ValidationError({'month': 'Must be 1-12.'})
    mine = m.CalendarEvent.objects.filter(Q(user__isnull=True) | Q(user=request.user))
    events = mine.filter(date__year=year, date__month=month)
    day = today if (today.year, today.month) == (year, month) else date(year, month, 1)
    day_events = mine.filter(date=day).exclude(time='')
    return Response({'year': year, 'monthIndex': month - 1, 'events': [{'id': event.id, 'date': str(event.date), 'title': event.title, 'type': event.type} for event in events],
                     'daySchedule': {'date': f'{cal.month_name[month]} {day.day}, {cal.day_name[day.weekday()]}',
                                     'items': [{'time': event.time, 'title': event.title} for event in day_events]}})


PROFILE_FIELDS = {'firstName': 'first_name', 'lastName': 'last_name', 'address': 'address', 'city': 'city', 'country': 'country',
                  'postalCode': 'postal_code'}
PROFILE_LIMITS = {'first_name': 150, 'last_name': 150, 'address': 255, 'city': 100, 'country': 100, 'postal_code': 20}


@api_view(['GET', 'PUT'])
def profile(request):
    user = request.user
    if request.method == 'PUT':
        data = request.data
        fields = []
        for js_name, model_field in PROFILE_FIELDS.items():
            if js_name in data:
                value = str(data.get(js_name) or '').strip()
                if len(value) > PROFILE_LIMITS[model_field]:
                    raise ValidationError({js_name: f'Keep this under {PROFILE_LIMITS[model_field]} characters.'})
                setattr(user, model_field, value)
                fields.append(model_field)
        if 'email' in data:
            email = str(data.get('email') or '').strip().lower()
            try:
                validate_email(email)
            except DjangoValidationError:
                raise ValidationError({'email': 'Enter a valid email address.'})
            if m.User.objects.exclude(pk=user.pk).filter(email__iexact=email).exists():
                raise ValidationError({'email': 'Another account already uses this email.'})
            user.email = email
            fields.append('email')
        if 'username' in data:
            username = str(data.get('username') or '').strip()
            if not re.fullmatch(r'[\w.@+-]{3,150}', username):
                raise ValidationError({'username': 'Use 3-150 letters, digits or . @ + - _ characters.'})
            if m.User.objects.exclude(pk=user.pk).filter(username__iexact=username).exists():
                raise ValidationError({'username': 'This username is taken.'})
            user.username = username
            fields.append('username')
        if fields:
            user.save(update_fields=fields)
    return Response(user_json(user, request))


@api_view(['POST'])
@parser_classes([MultiPartParser])
def profile_avatar(request):
    upload = request.FILES.get('avatar')
    if not upload:
        raise ValidationError({'avatar': 'No file uploaded.'})
    if upload.size > 3 * 1024 * 1024 or not str(upload.content_type or '').startswith('image/'):
        raise ValidationError({'avatar': 'Use an image smaller than 3 MB.'})
    request.user.avatar = upload
    request.user.save(update_fields=['avatar'])
    return Response({'url': abs_url(request, request.user.avatar)})
